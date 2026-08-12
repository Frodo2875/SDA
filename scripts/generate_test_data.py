"""Generate fixed virtual student data for local development and testing."""

from collections import Counter
from pathlib import Path

from docx import Document
from openpyxl import Workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"

STUDENTS = [
    ("S001", "张三", "计算机科学", "研一", "1班"),
    ("S002", "李四", "计算机科学", "研一", "1班"),
    ("S003", "王五", "软件工程", "研二", "2班"),
    ("S004", "张三", "人工智能", "研一", "3班"),
    ("S005", "赵六", "软件工程", "研一", "2班"),
    ("S006", "钱七", "人工智能", "研二", "3班"),
    ("S007", "孙八", "数据科学", "研一", "4班"),
    ("S008", "周九", "计算机科学", "研二", "1班"),
    ("S009", "吴十", "数据科学", "研二", "4班"),
    ("S010", "郑晓雨", "网络空间安全", "研一", "5班"),
    ("S011", "冯晨", "网络空间安全", "研二", "5班"),
    ("S012", "陈星", "软件工程", "研一", "2班"),
]

# S012 intentionally has no score record.
SCORES = [
    ("S001", 88, 91, 93),
    ("S002", 82, 86, 89),
    ("S003", 90, 84, 92),
    ("S004", 95, 89, 96),
    ("S005", 78, 83, 85),
    ("S006", 87, 92, 90),
    ("S007", 91, 88, 94),
    ("S008", 85, 90, 87),
    ("S009", 84, 86, 88),
    ("S010", 89, 85, 91),
    ("S011", 81, 88, 86),
]

# S011 intentionally has no research record.
RESEARCH = [
    ("S001", 1, 0, 2),
    ("S002", 0, 1, 1),
    ("S003", 2, 0, 1),
    ("S004", 1, 1, 3),
    ("S005", 0, 0, 1),
    ("S006", 2, 1, 0),
    ("S007", 1, 0, 2),
    ("S008", 0, 1, 2),
    ("S009", 1, 0, 1),
    ("S010", 0, 0, 2),
    ("S012", 1, 1, 0),
]


def write_workbook(path: Path, headers: list[str], rows: list[tuple]) -> None:
    """Write rows to a single-sheet Excel workbook."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


def build_score_rows() -> list[tuple]:
    """Calculate averages and rankings within each major."""
    major_by_student = {student_id: major for student_id, _, major, _, _ in STUDENTS}
    score_details = []

    for student_id, math, english, major_course in SCORES:
        average = round((math + english + major_course) / 3, 2)
        score_details.append(
            {
                "student_id": student_id,
                "math": math,
                "english": english,
                "major_course": major_course,
                "average": average,
                "major": major_by_student[student_id],
            }
        )

    rankings = {}
    majors = {detail["major"] for detail in score_details}
    for major in majors:
        major_scores = sorted(
            (detail for detail in score_details if detail["major"] == major),
            key=lambda detail: detail["average"],
            reverse=True,
        )
        for rank, detail in enumerate(major_scores, start=1):
            rankings[detail["student_id"]] = rank

    return [
        (
            detail["student_id"],
            detail["math"],
            detail["english"],
            detail["major_course"],
            detail["average"],
            rankings[detail["student_id"]],
        )
        for detail in score_details
    ]


def create_word_template(path: Path) -> None:
    """Create a Word template containing only its title."""
    document = Document()
    document.add_heading("学生综合评价", level=1)
    document.save(path)


def format_students(student_ids: list[str]) -> str:
    """Format student IDs and names for the generation summary."""
    name_by_id = {student_id: name for student_id, name, _, _, _ in STUDENTS}
    return "、".join(f"{student_id} {name_by_id[student_id]}" for student_id in student_ids)


def main() -> None:
    """Generate all fixed test files and print a data summary."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    student_path = DATA_DIR / "学生基本信息.xlsx"
    score_path = DATA_DIR / "学生成绩.xlsx"
    research_path = DATA_DIR / "科研成果.xlsx"
    evaluation_path = DATA_DIR / "综合评价.docx"

    score_rows = build_score_rows()
    write_workbook(student_path, ["学号", "姓名", "专业", "年级", "班级"], STUDENTS)
    write_workbook(
        score_path,
        ["学号", "数学", "英语", "专业课", "平均分", "专业排名"],
        score_rows,
    )
    write_workbook(research_path, ["学号", "论文数", "专利数", "竞赛数"], RESEARCH)
    create_word_template(evaluation_path)

    all_student_ids = {student[0] for student in STUDENTS}
    score_student_ids = {score[0] for score in SCORES}
    research_student_ids = {record[0] for record in RESEARCH}
    name_counts = Counter(student[1] for student in STUDENTS)
    duplicate_names = sorted(name for name, count in name_counts.items() if count > 1)
    duplicate_students = [
        student[0] for student in STUDENTS if student[1] in duplicate_names
    ]
    missing_scores = sorted(all_student_ids - score_student_ids)
    missing_research = sorted(all_student_ids - research_student_ids)

    print("创建的文件：")
    for path in (student_path, score_path, research_path, evaluation_path):
        print(f"- {path.relative_to(PROJECT_ROOT)}")
    print(f"学生数量：{len(STUDENTS)}")
    print(f"成绩记录数量：{len(score_rows)}")
    print(f"科研记录数量：{len(RESEARCH)}")
    print(f"重名学生：{format_students(duplicate_students)}")
    print(f"缺失成绩学生：{format_students(missing_scores)}")
    print(f"缺失科研学生：{format_students(missing_research)}")


if __name__ == "__main__":
    main()
