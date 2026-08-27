"""V3.5 DocumentBlock title, paragraph, table and legacy compatibility tests."""

from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

from backend import database
from backend.document_blocks import DocumentBlock, blocks_from_chunks
from backend.services.document_index import index_document, retrieve_document
from backend.table_structure import TableStructure
from backend.tools import excel_utils


@pytest.fixture
def indexed_word(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> dict:
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    path = upload_dir / "结构化材料.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "学生材料页眉"
    document.sections[0].footer.paragraphs[0].text = "学生材料页脚"
    document.add_heading("材料提交说明", level=1)
    document.add_paragraph("正文规定材料提交截止日期为六月三十日。")
    image_path = tmp_path / "layout-image.png"
    Image.new("RGB", (20, 20), "white").save(image_path)
    document.add_picture(str(image_path), width=Inches(0.2))
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "材料类型"
    table.cell(0, 1).text = "份数"
    table.cell(1, 0).text = "成绩单"
    table.cell(1, 1).text = "一份"
    document.save(path)
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)
    database.register_file(
        file_name=path.name,
        file_type="word",
        file_path=f"data/uploads/{path.name}",
        lifecycle_status="ready",
        parse_status="parsed",
        queryable=True,
    )
    record = database.get_file_record(path.name)
    result = index_document(record["file_id"])
    assert result["ok"] is True
    return {"record": record, "result": result}


def test_title_is_persisted_as_document_block(indexed_word: dict) -> None:
    blocks = [DocumentBlock.model_validate(item) for item in indexed_word["result"]["data"]["blocks"]]
    title = next(block for block in blocks if block.block_type == "title")

    assert title.content == "材料提交说明"
    assert title.file_id == indexed_word["record"]["file_id"]
    assert title.page_no is None
    assert title.bbox is None
    assert title.confidence is None
    assert title.source_parser == "python-docx"
    assert title.parent_id is None
    chunks = database.get_document_chunks(title.file_id)
    title_chunk = next(chunk for chunk in chunks if chunk["metadata"].get("block_id") == title.block_id)
    assert title_chunk["chunk_text"] == title.content
    assert title_chunk["metadata"]["block_type"] == "title"


def test_paragraph_chunk_produces_block_referenced_evidence(indexed_word: dict) -> None:
    file_id = indexed_word["record"]["file_id"]
    chunks = database.get_document_chunks(file_id)
    paragraph_chunk = next(
        chunk for chunk in chunks if chunk["metadata"].get("block_type") == "paragraph"
    )

    retrieved = retrieve_document({"file_id": file_id}, "材料提交截止日期")

    assert retrieved["ok"] is True
    evidence = retrieved["data"]["evidence"][0]
    assert evidence["chunk_id"] == paragraph_chunk["chunk_id"]
    assert evidence["block_id"] == paragraph_chunk["metadata"]["block_id"]
    assert paragraph_chunk["metadata"]["block"]["content"] == (
        "正文规定材料提交截止日期为六月三十日。"
    )
    blocks = [DocumentBlock.model_validate(item) for item in indexed_word["result"]["data"]["blocks"]]
    title = next(block for block in blocks if block.block_type == "title")
    paragraph = next(block for block in blocks if block.block_type == "paragraph")
    assert paragraph.parent_id == title.block_id
    assert paragraph.source_parser == "python-docx"
    assert paragraph.bbox is None
    assert paragraph.confidence is None


def test_o06_header_footer_and_image_blocks_use_existing_schema(indexed_word: dict) -> None:
    blocks = [DocumentBlock.model_validate(item) for item in indexed_word["result"]["data"]["blocks"]]
    by_type = {block.block_type: block for block in blocks}
    assert by_type["header"].content == "学生材料页眉"
    assert by_type["footer"].content == "学生材料页脚"
    assert by_type["image"].content == ""
    assert by_type["image"].status == "degraded"
    assert by_type["image"].warnings == ["IMAGE_TEXT_NOT_EXTRACTED"]
    assert all(by_type[name].source_parser == "python-docx" for name in ("header", "footer", "image"))


def test_table_and_cells_are_distinct_blocks_and_generate_chunks(
    indexed_word: dict,
) -> None:
    blocks = [DocumentBlock.model_validate(item) for item in indexed_word["result"]["data"]["blocks"]]
    table = next(block for block in blocks if block.block_type == "table")
    cells = [block for block in blocks if block.block_type == "cell"]

    assert table.content == "材料类型\t份数\n成绩单\t一份"
    assert [cell.content for cell in cells] == ["材料类型", "份数", "成绩单", "一份"]
    chunks = database.get_document_chunks(indexed_word["record"]["file_id"])
    assert {chunk["metadata"]["block_type"] for chunk in chunks} >= {
        "table",
        "cell",
    }
    assert all(
        any(chunk["metadata"].get("block_id") == block.block_id for chunk in chunks)
        for block in [table, *cells]
    )
    table_chunk = next(
        chunk for chunk in chunks if chunk["metadata"].get("block_id") == table.block_id
    )
    structure = TableStructure.model_validate(table_chunk["metadata"]["table_structure"])
    assert structure.table_id == table.block_id
    assert structure.row_count == 2
    assert structure.column_count == 2
    assert len(structure.rows) == 2
    assert len(structure.columns) == 2
    assert len(structure.cells) == 4
    assert [(cell.row_index, cell.column_index, cell.cell_text) for cell in structure.cells] == [
        (0, 0, "材料类型"), (0, 1, "份数"), (1, 0, "成绩单"), (1, 1, "一份")
    ]
    assert all(cell.bbox is None and cell.confidence is None for cell in structure.cells)
    assert all(cell.parent_id == table.block_id for cell in cells)
    cell_chunks = [chunk for chunk in chunks if chunk["metadata"].get("cell_id")]
    assert {
        "table_id", "cell_id", "row_index", "column_index", "bbox", "confidence"
    } <= set(cell_chunks[0]["metadata"])

    retrieved = retrieve_document(
        {"file_id": indexed_word["record"]["file_id"]}, "成绩单", top_k=10
    )
    locators = [
        (evidence.get("table"), evidence.get("cell"))
        for evidence in retrieved["evidence"]
    ]
    expected_cell = next(cell for cell in structure.cells if cell.cell_text == "成绩单")
    assert (table.block_id, expected_cell.cell_id) in locators


def test_o08_merged_table_degrades_without_inventing_cells(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    path = upload_dir / "复杂表格.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "合并标题"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = "原始内容甲"
    table.cell(1, 1).text = "原始内容乙"
    document.save(path)
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)
    database.register_file(
        file_name=path.name, file_type="word", file_path=f"data/uploads/{path.name}",
        lifecycle_status="ready", parse_status="parsed", queryable=True,
    )
    record = database.get_file_record(path.name)
    result = index_document(record["file_id"])
    assert result["ok"] is True
    blocks = [DocumentBlock.model_validate(item) for item in result["data"]["blocks"]]
    table_block = next(block for block in blocks if block.block_type == "table")
    assert table_block.status == "degraded"
    assert table_block.warnings == ["MERGED_OR_IRREGULAR_CELLS_UNSUPPORTED"]
    assert [block for block in blocks if block.block_type == "cell"] == []
    chunks = database.get_document_chunks(record["file_id"])
    table_chunk = next(chunk for chunk in chunks if chunk["metadata"].get("block_type") == "table")
    structure = TableStructure.model_validate(table_chunk["metadata"]["table_structure"])
    assert structure.status == "degraded"
    assert structure.rows == structure.columns == structure.cells == []
    assert "原始内容甲" in structure.raw_text
    retrieved = retrieve_document({"file_id": record["file_id"]}, "原始内容甲", 5)
    assert retrieved["data"]["status"] == "found"


def test_legacy_chunk_without_block_metadata_remains_readable() -> None:
    legacy = {
        "chunk_id": "legacy-chunk",
        "file_id": "legacy-file",
        "page_no": 1,
        "chunk_index": 0,
        "chunk_text": "旧数据",
        "text_hash": "legacy-hash",
        "metadata": {"source_type": "pdf", "page_no": 1},
    }

    assert blocks_from_chunks([legacy]) == []
    assert legacy["metadata"] == {"source_type": "pdf", "page_no": 1}
