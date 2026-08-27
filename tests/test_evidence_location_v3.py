"""V3.19 R206-R209 Evidence source-navigation tests."""

from pathlib import Path

from docx import Document
from openpyxl import Workbook

from backend import database
from backend.agent import _collect_evidence
from backend.document_blocks import block_to_chunks, make_document_block
from backend.evidence import build_evidence, make_evidence_id
from backend.services.evidence_locator import locate_evidence
from backend.tools import excel_utils
from backend.tools.schema_tools import inspect_excel
from backend.tools.table_tools import query_table


def _register_document(
    tmp_path: Path, *, file_name: str, file_type: str
) -> tuple[str, Path]:
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir(exist_ok=True)
    path = upload_dir / file_name
    path.write_bytes(b"source")
    database.register_file(
        file_name=file_name,
        file_type=file_type,
        file_path=f"data/uploads/{file_name}",
        lifecycle_status="ready",
        parse_status="parsed",
        queryable=True,
        index_status="indexed",
    )
    return database.get_file_record(file_name)["file_id"], path


def test_r206_pdf_location_returns_correct_page_bbox_and_confidence(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)
    file_id, _ = _register_document(
        tmp_path, file_name="扫描材料.pdf", file_type="pdf"
    )
    block = make_document_block(
        file_id=file_id,
        sequence=0,
        block_type="paragraph",
        content="第七页奖学金申请原文。",
        page_no=7,
        bbox=[10.0, 20.0, 200.0, 80.0],
        confidence=0.78,
        source_parser="rapidocr",
    )
    chunk = block_to_chunks(block, source_type="ocr")[0]
    database.replace_document_chunks(file_id, [chunk])
    evidence = build_evidence(
        evidence_id=make_evidence_id(file_id=file_id, chunk_id=chunk["chunk_id"]),
        source_type="unstructured",
        file_id=file_id,
        file_name="扫描材料.pdf",
        page_no=7,
        chunk_id=chunk["chunk_id"],
        block_id=block.block_id,
        bbox=block.bbox,
        confidence=block.confidence,
        value_summary=block.content,
    )

    result = locate_evidence(evidence["evidence_id"], session_id="r206")

    assert result["ok"] is True
    location = result["data"]
    assert location["location_type"] == "pdf"
    assert location["page_no"] == 7
    assert location["block_id"] == block.block_id
    assert location["bbox"] == [10.0, 20.0, 200.0, 80.0]
    assert location["highlight"] == {"bbox": [10.0, 20.0, 200.0, 80.0]}
    assert location["confidence"] == 0.78
    assert location["text"] == block.content


def test_r207_excel_location_opens_exact_sheet_row_and_field(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)
    path = tmp_path / "学生成绩.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "成绩表"
    worksheet.append(["学号", "平均成绩"])
    worksheet.append(["S001", 92])
    worksheet.append(["S002", 88])
    workbook.save(path)
    workbook.close()
    database.initialize_database()
    record = database.get_file_record(path.name)
    assert inspect_excel(record["file_id"])["ok"] is True
    query = query_table(
        record["file_id"],
        "成绩表",
        filters=[{"field": "学号", "op": "=", "value": "S001"}],
        select=["平均成绩"],
    )
    evidence = query["evidence_chain"][0]

    result = locate_evidence(evidence["evidence_id"], session_id="r207")

    assert result["ok"] is True
    location = result["data"]
    assert location["location_type"] == "excel"
    assert location["sheet"] == "成绩表"
    assert location["table_id"] == "成绩表"
    assert location["cell_id"] == "B2"
    assert location["row_index"] == 2
    assert location["column_index"] == 2
    assert location["record_identifier"] == "S001"
    assert location["field"] == "平均成绩"
    assert location["target_value"] == 92
    assert location["row_values"] == ["S001", 92]


def test_r208_word_location_opens_indexed_paragraph(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)
    file_id, path = _register_document(
        tmp_path, file_name="申请说明.docx", file_type="word"
    )
    document = Document()
    document.add_paragraph("第一段说明。")
    document.add_paragraph("需要定位的第二段原文。")
    document.save(path)
    block = make_document_block(
        file_id=file_id,
        sequence=1,
        block_type="paragraph",
        content="需要定位的第二段原文。",
        source_parser="python-docx",
    )
    chunk = block_to_chunks(
        block,
        source_type="word",
        metadata={"paragraph_no": 2, "paragraph_start": 2},
    )[0]
    database.replace_document_chunks(file_id, [chunk])
    evidence = build_evidence(
        evidence_id=make_evidence_id(file_id=file_id, block_id=block.block_id),
        source_type="unstructured",
        file_id=file_id,
        file_name="申请说明.docx",
        chunk_id=chunk["chunk_id"],
        block_id=block.block_id,
        value_summary=block.content,
    )

    result = locate_evidence(evidence["evidence_id"], session_id="r208")

    assert result["ok"] is True
    assert result["data"]["location_type"] == "word"
    assert result["data"]["paragraph_no"] == 2
    assert result["data"]["block_id"] == block.block_id
    assert result["data"]["text"] == block.content


def test_r209_location_failure_keeps_safe_message_and_records_trace() -> None:
    result = locate_evidence("f" * 24, session_id="r209")

    assert result["ok"] is False
    assert result["message"] == "来源存在，但当前无法打开对应预览位置"
    traces = database.get_session_trace_records("r209")
    assert len(traces) == 1
    assert traces[0]["event_type"] == "evidence_location"
    assert traces[0]["tool_name"] == "locate_evidence"
    assert traces[0]["result_status"] == "failed"
    assert traces[0]["error_code"] == "EVIDENCE_LOCATION_NOT_FOUND"


def test_unused_retrieval_source_is_not_exposed_after_evaluation() -> None:
    used = {"evidence_id": "used", "file_id": "file-used"}
    unused = {"evidence_id": "unused", "file_id": "file-unused"}
    calls = [
        {"name": "retrieve_document", "result": {"evidence": [used, unused]}},
        {
            "name": "evaluate_scholarship_eligibility",
            "result": {
                "data": {
                    "evidence_chain": [used],
                    "used_tools": ["retrieve_document"],
                }
            },
        },
    ]

    assert _collect_evidence(calls) == [used]
