"""V2.10 B01-B05 Batch filtering, progress, partial failure, and HITL tests."""

import shutil
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from backend import database
from backend.services import batch_service
from backend.services.batch_service import get_batch, run_batch
from backend.services.confirmation import confirm_action
from backend.tools import excel_utils
from backend.tools.schema_tools import inspect_excel


PROJECT_ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def low_score_materials(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    for source in (PROJECT_ROOT / "data").iterdir():
        if source.is_file() and source.suffix.lower() in {".xlsx", ".docx"}:
            shutil.copy2(source, tmp_path / source.name)
    workbook = load_workbook(tmp_path / "学生成绩.xlsx")
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    for row in range(2, sheet.max_row + 1):
        student_id = str(sheet.cell(row, headers["学号"]).value)
        if student_id == "S001":
            values = (60, 60, 60)
        elif student_id == "S011":
            values = (65, 65, 65)
        else:
            continue
        for field, value in zip(("数学", "英语", "专业课"), values, strict=True):
            sheet.cell(row, headers[field], value)
    workbook.save(tmp_path / "学生成绩.xlsx")
    workbook.close()
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)


def _successful_items(result: dict) -> list[dict]:
    return [item for item in result["data"]["items"] if item["status"] == "success"]


@pytest.mark.anyio
async def test_b01_python_filters_students_below_70(low_score_materials) -> None:
    result = await run_batch(
        {
            "session_id": "b01",
            "action_type": "score_below_threshold",
            "threshold": 70,
        }
    )

    assert result["ok"] is True
    assert result["data"]["status"] == "success"
    assert {item["target_id"] for item in _successful_items(result)} == {"S001", "S011"}
    assert result["data"]["summary"] == {
        "total": 12, "success": 2, "failed": 1, "skipped": 9
    }
    assert result["data"]["failure_details"][0]["target_id"] == "S012"


@pytest.mark.anyio
async def test_b02_below_70_and_no_research_preserves_no_record_semantics(
    low_score_materials,
) -> None:
    result = await run_batch(
        {
            "session_id": "b02",
            "action_type": "score_below_threshold_no_research",
            "threshold": 70,
        }
    )

    assert result["ok"] is True
    successful = _successful_items(result)
    assert [item["target_id"] for item in successful] == ["S011"]
    assert "未查询到相关记录" in successful[0]["result_summary"]
    assert "没有科研成果" not in successful[0]["result_summary"]
    assert result["data"]["summary"] == {
        "total": 12, "success": 1, "failed": 1, "skipped": 10
    }


@pytest.mark.anyio
async def test_b03_college_count_uses_python_aggregate(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)
    upload = tmp_path / "uploads"
    upload.mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "学生"
    sheet.append(["学号", "姓名", "学院"])
    sheet.append(["C001", "甲", "计算机学院"])
    sheet.append(["C002", "乙", "计算机学院"])
    sheet.append(["C003", "丙", "软件学院"])
    workbook.save(upload / "学院统计.xlsx")
    workbook.close()
    database.register_file(
        file_name="学院统计.xlsx", file_type="excel",
        file_path="data/uploads/学院统计.xlsx", lifecycle_status="ready",
        parse_status="parsed", queryable=True,
    )
    record = database.get_file_record("学院统计.xlsx")
    assert inspect_excel(record["file_id"])["ok"] is True

    result = await run_batch(
        {
            "session_id": "b03", "action_type": "college_count",
            "file_id": record["file_id"], "sheet": "学生",
            "college_field": "学院",
        }
    )

    assert result["ok"] is True
    counts = {
        item["target_id"]: int(item["result_summary"].split("：")[1])
        for item in result["data"]["items"]
    }
    assert counts == {"计算机学院": 2, "软件学院": 1}
    assert result["data"]["summary"] == {
        "total": 2, "success": 2, "failed": 0, "skipped": 0
    }


@pytest.mark.anyio
async def test_b04_quality_batch_continues_after_one_item_failure(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def quality(student_id: str) -> dict:
        if student_id == "Q002":
            return {
                "ok": False, "data": None, "error_code": "TEMPORARY_IO_ERROR",
                "message": "模拟单项读取失败",
            }
        return {
            "ok": True,
            "data": {"status": "found", "student_id": student_id, "issues": []},
            "error_code": None,
            "message": "完成",
        }

    monkeypatch.setattr(batch_service, "validate_student_data", quality)
    result = await run_batch(
        {
            "session_id": "b04", "action_type": "data_quality",
            "student_ids": ["Q001", "Q002", "Q003"],
        }
    )

    assert result["ok"] is True
    assert result["data"]["status"] == "success"
    assert result["data"]["summary"] == {
        "total": 3, "success": 2, "failed": 1, "skipped": 0
    }
    assert result["data"]["failure_details"] == [
        {
            "target_id": "Q002", "error_code": "TEMPORARY_IO_ERROR",
            "message": "模拟单项读取失败",
        }
    ]


@pytest.mark.anyio
async def test_b05_twenty_evaluations_keep_18_successes_and_use_one_hitl_action(
    low_score_materials,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    student_ids = [f"B{index:03d}" for index in range(1, 21)]
    target = excel_utils.DATA_DIR / "综合评价.docx"
    original = target.read_bytes()
    generated_ids: list[str] = []

    monkeypatch.setattr(
        batch_service,
        "get_student_info",
        lambda student_id: {
            "ok": True,
            "data": {"student_id": student_id, "name": f"虚拟{student_id}"},
            "error_code": None,
            "message": "完成",
        },
    )
    monkeypatch.setattr(
        batch_service,
        "get_student_scores",
        lambda student_id: {
            "ok": True,
            "data": (
                {"status": "no_record", "student_id": student_id}
                if student_id in {"B019", "B020"}
                else {"status": "found", "student_id": student_id, "average_score": 80}
            ),
            "error_code": None,
            "message": "完成",
        },
    )
    monkeypatch.setattr(
        batch_service,
        "get_student_research",
        lambda student_id: {
            "ok": True,
            "data": {"status": "no_record", "student_id": student_id},
            "error_code": None,
            "message": "当前科研成果材料中未查询到相关记录。",
        },
    )

    async def generator(context: dict) -> str:
        student_id = context["student"]["student_id"]
        generated_ids.append(student_id)
        return f"{student_id} 学习表现稳定；当前科研成果材料中未查询到相关记录。"

    result = await run_batch(
        {
            "session_id": "b05", "action_type": "generate_evaluations",
            "student_ids": student_ids, "target_file": "综合评价.docx",
        },
        evaluation_generator=generator,
    )

    assert result["ok"] is True
    assert result["data"]["status"] == "pending"
    assert result["data"]["summary"] == {
        "total": 20, "success": 18, "failed": 2, "skipped": 0
    }
    assert generated_ids == student_ids[:18]
    assert target.read_bytes() == original
    action = result["data"]["pending_action"]
    assert action["status"] == "pending"
    assert action["diff_preview"]["operation_type"] == "append"
    assert {item["action_id"] for item in result["data"]["items"] if item["status"] == "success"} == {action["action_id"]}

    first = confirm_action(action["action_id"])
    bytes_after_first = target.read_bytes()
    second = confirm_action(action["action_id"])
    completed = get_batch(result["data"]["batch_id"])

    assert first["ok"] is True
    assert second["error_code"] == "ACTION_NOT_PENDING"
    assert target.read_bytes() == bytes_after_first
    assert completed["data"]["status"] == "success"
    assert completed["data"]["summary"] == {
        "total": 20, "success": 18, "failed": 2, "skipped": 0
    }
    paragraphs = [paragraph.text for paragraph in Document(target).paragraphs]
    assert sum("B001 综合评价" in paragraph for paragraph in paragraphs) == 1
