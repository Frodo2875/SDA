"""Automated tests for the deterministic Python tools."""

import shutil
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from backend.tools import excel_utils
from backend.tools.analysis_tools import compare_students
from backend.tools.file_tools import get_file_info, list_files
from backend.tools.student_tools import (
    get_student_info,
    get_student_research,
    get_student_scores,
    search_student,
)
from backend.tools.word_tools import read_word, write_word


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = PROJECT_ROOT / "data"
EXPECTED_FILES = {
    "学生基本信息.xlsx",
    "学生成绩.xlsx",
    "科研成果.xlsx",
    "综合评价.docx",
}


def test_list_files_returns_four_test_files() -> None:
    result = list_files()

    assert result["ok"] is True
    assert result["error_code"] is None
    assert {item["file_name"] for item in result["data"]} == EXPECTED_FILES
    assert len(result["data"]) == 4
    assert all(item["exists"] and item["size"] > 0 for item in result["data"])


def test_search_student_by_id_returns_unique_match() -> None:
    result = search_student("S001")

    assert result["ok"] is True
    assert result["data"]["status"] == "found"
    assert result["data"]["student"] == {
        "student_id": "S001",
        "name": "张三",
        "major": "计算机科学",
        "grade": "研一",
        "class_name": "1班",
    }


def test_search_duplicate_name_returns_ambiguous_candidates() -> None:
    result = search_student("张三")

    assert result["ok"] is True
    assert result["data"]["status"] == "ambiguous"
    assert {item["student_id"] for item in result["data"]["candidates"]} == {
        "S001",
        "S004",
    }


def test_search_unknown_student_returns_not_found() -> None:
    result = search_student("不存在学生")

    assert result["ok"] is True
    assert result["data"] == {"status": "not_found"}


def test_get_student_info_returns_expected_data() -> None:
    result = get_student_info("S003")

    assert result["ok"] is True
    assert result["data"] == {
        "student_id": "S003",
        "name": "王五",
        "major": "软件工程",
        "grade": "研二",
        "class_name": "2班",
    }


def test_get_student_scores_returns_expected_data() -> None:
    result = get_student_scores("S001")

    assert result["ok"] is True
    assert result["data"] == {
        "status": "found",
        "student_id": "S001",
        "scores": {"数学": 88, "英语": 91, "专业课": 93},
        "average_score": 90.67,
        "rank": 1,
    }


def test_calculated_average_matches_excel_source_values() -> None:
    workbook = load_workbook(DATA_DIR / "学生成绩.xlsx", read_only=True, data_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    headers = rows[0]
    source = dict(zip(headers, next(row for row in rows[1:] if str(row[0]) == "S001")))
    expected_average = round((source["数学"] + source["英语"] + source["专业课"]) / 3, 2)

    result = get_student_scores("S001")

    assert result["data"]["average_score"] == expected_average
    assert source["平均分"] == expected_average


def test_get_student_research_returns_expected_data() -> None:
    result = get_student_research("S001")

    assert result["ok"] is True
    assert result["data"] == {
        "status": "found",
        "student_id": "S001",
        "research": {"论文数": 1, "专利数": 0, "竞赛数": 2},
    }


def test_missing_research_record_returns_no_record() -> None:
    result = get_student_research("S011")

    assert result["ok"] is True
    assert result["data"] == {"status": "no_record", "student_id": "S011"}
    assert "不能据此认定科研成果为零" in result["message"]


def test_missing_score_record_returns_no_record() -> None:
    result = get_student_scores("S012")

    assert result["ok"] is True
    assert result["data"] == {"status": "no_record", "student_id": "S012"}


def test_compare_students_calculates_differences() -> None:
    result = compare_students(["S001", "S002"])

    assert result["ok"] is True
    assert result["data"]["differences"] == [
        {
            "left_student_id": "S001",
            "right_student_id": "S002",
            "calculation": "left_minus_right",
            "values": {
                "average_score": 5.0,
                "rank": -2,
                "paper_count": 1,
                "patent_count": -1,
                "competition_count": 1,
            },
        }
    ]


def test_read_word_returns_template_title() -> None:
    result = read_word("综合评价.docx")

    assert result["ok"] is True
    assert result["data"]["text"] == "学生综合评价"
    assert result["data"]["paragraphs"] == ["学生综合评价"]


def test_write_word_appends_to_isolated_copy(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    isolated_document = tmp_path / "综合评价.docx"
    shutil.copy2(DATA_DIR / "综合评价.docx", isolated_document)
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)

    result = write_word("综合评价.docx", "这是一条隔离的测试评价。")
    document = Document(isolated_document)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]

    assert result["ok"] is True
    assert result["data"]["mode"] == "append"
    assert paragraphs == ["学生综合评价", "这是一条隔离的测试评价。"]


def test_word_write_does_not_modify_official_document(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    official_document = DATA_DIR / "综合评价.docx"
    original_bytes = official_document.read_bytes()
    isolated_document = tmp_path / "综合评价.docx"
    shutil.copy2(official_document, isolated_document)
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)

    result = write_word("综合评价.docx", "不会进入正式文件的内容。")

    assert result["ok"] is True
    assert official_document.read_bytes() == original_bytes
    assert "不会进入正式文件的内容。" in read_word_from_path(isolated_document)


def read_word_from_path(path: Path) -> list[str]:
    """Read paragraphs directly for assertions on an isolated document."""
    return [paragraph.text for paragraph in Document(path).paragraphs if paragraph.text]


@pytest.mark.parametrize(
    ("call", "expected_error"),
    [
        (lambda: get_file_info("不存在.xlsx"), "FILE_NOT_FOUND"),
        (lambda: read_word("不存在.docx"), "FILE_NOT_FOUND"),
        (lambda: write_word("不存在.docx", "内容"), "FILE_NOT_FOUND"),
    ],
)
def test_missing_file_returns_clear_error(call, expected_error: str) -> None:
    result = call()

    assert result["ok"] is False
    assert result["error_code"] == expected_error
    assert "不存在" in result["message"]


def test_excel_missing_required_field_returns_clear_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["学号", "姓名", "专业", "年级"])
    worksheet.append(["T001", "测试学生", "测试专业", "研一"])
    workbook.save(tmp_path / "学生基本信息.xlsx")
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)

    result = get_student_info("T001")

    assert result["ok"] is False
    assert result["error_code"] == "MISSING_FIELDS"
    assert result["data"] == {"missing_fields": ["班级"]}
    assert "班级" in result["message"]
