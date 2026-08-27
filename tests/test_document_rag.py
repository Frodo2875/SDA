"""V2.6 R01/R02/R04/R05 PDF, Word chunks, retrieval, and cleanup tests."""

import sqlite3
from pathlib import Path

import pytest
from docx import Document

from backend import agent, database
from backend.services import document_index
from backend.services.document_index import (
    index_document,
    parse_pdf,
    reindex_document,
    reprocess_document,
    retrieve_document,
)
from backend.services.file_lifecycle import delete_uploaded_file, get_file_lifecycle
from backend.services.file_upload import save_uploaded_file
from backend.tools import excel_utils
from backend.tools.file_tools import list_files


@pytest.fixture
def document_data_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)
    (tmp_path / "uploads").mkdir()
    return tmp_path


def _register(
    data_dir: Path,
    file_name: str,
    file_type: str,
    content: bytes = b"placeholder",
) -> tuple[str, Path]:
    path = data_dir / "uploads" / file_name
    path.write_bytes(content)
    database.register_file(
        file_name=file_name,
        file_type=file_type,
        file_path=f"data/uploads/{file_name}",
        lifecycle_status="ready",
        parse_status="parsed",
        queryable=True,
    )
    record = database.get_file_record(file_name)
    assert record is not None
    return record["file_id"], path


def test_r01_text_pdf_preserves_each_page_and_returns_page_evidence(
    document_data_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    file_id, _ = _register(document_data_dir, "培养规定.pdf", "pdf")
    monkeypatch.setattr(
        document_index,
        "_load_pdf_pages",
        lambda path: [
            (1, "第一页介绍培养目标。"),
            (2, "第二页规定申请奖学金需要提交成绩证明。"),
            (3, ""),
        ],
    )

    parsed = parse_pdf(file_id)
    chunks = database.get_document_chunks(file_id)
    database.initialize_database()
    retrieved = retrieve_document(
        {"file_ids": [file_id]}, "申请奖学金需要提交成绩证明", top_k=5
    )

    assert parsed["ok"] is True
    assert parsed["data"]["page_count"] == 3
    assert parsed["data"]["chunk_count"] == 3
    assert [item["page_no"] for item in chunks] == [1, 2, 3]
    assert chunks[2]["chunk_text"] == ""
    assert retrieved["data"]["status"] == "found"
    assert retrieved["evidence"] == retrieved["data"]["evidence"]
    evidence = retrieved["data"]["evidence"][0]
    assert evidence["page_no"] == 2
    assert {
        "evidence_id",
        "task_id",
        "source_type",
        "file_id",
        "file_name",
        "sheet",
        "page_no",
        "chunk_id",
        "block_id",
        "field",
        "record_key",
        "value_summary",
        "score",
        "text_excerpt",
    } == set(evidence)
    assert evidence["source_type"] == "unstructured"
    assert evidence["block_id"] == chunks[1]["metadata"]["block_id"]
    assert "申请奖学金需要提交成绩证明" in evidence["text_excerpt"]


def test_r02_scanned_pdf_returns_explicit_ocr_message_and_no_chunks(
    document_data_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    file_id, _ = _register(document_data_dir, "扫描材料.pdf", "pdf")
    monkeypatch.setattr(
        document_index,
        "_load_pdf_pages",
        lambda path: [(1, ""), (2, "   ")],
    )
    monkeypatch.setattr(
        document_index.ocr_service,
        "ocr_pdf",
        lambda path: {
            "ok": False,
            "data": None,
            "error_code": "OCR_PROCESSING_ERROR",
            "message": "OCR 识别过程失败",
        },
    )

    result = parse_pdf(file_id)

    assert result["ok"] is False
    assert result["error_code"] == "OCR_PROCESSING_ERROR"
    assert result["message"] == "OCR 识别过程失败"
    assert database.get_document_chunks(file_id) == []
    record = database.get_file_record_by_id(file_id)
    assert record["parse_status"] == "failed"
    assert record["index_status"] == "failed"
    assert record["queryable"] == 0


def test_pdf_upload_is_supported_and_automatically_indexed(
    document_data_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        document_index,
        "_load_pdf_pages",
        lambda path: [(1, "上传 PDF 的有效文本内容。")],
    )

    result = save_uploaded_file("上传文本.pdf", b"%PDF test fixture")

    assert result["ok"] is True
    assert result["data"]["file_type"] == "pdf"
    assert result["data"]["index_status"] == "indexed"
    assert database.get_document_chunks(result["data"]["file_id"])[0]["page_no"] == 1
    listed = next(item for item in list_files()["data"] if item["file_name"] == "上传文本.pdf")
    assert listed["file_type"] == "pdf"
    assert listed["index_status"] == "indexed"


def test_r04_long_word_retrieves_only_the_relevant_chunk(
    document_data_dir: Path,
) -> None:
    file_id, path = _register(document_data_dir, "长篇制度.docx", "word", b"")
    document = Document()
    document.add_heading("第一章 无关说明", level=1)
    document.add_paragraph("甲" * 830)
    document.add_heading("第二章 请假规则", level=1)
    document.add_paragraph("研究生连续请假超过三天需要导师签字并提交学院审批。")
    document.add_heading("第三章 其他说明", level=1)
    document.add_paragraph("乙" * 830)
    document.save(path)

    indexed = index_document(file_id)
    chunks = database.get_document_chunks(file_id)
    result = retrieve_document(
        {"file_ids": [file_id]}, "连续请假超过三天需要导师签字", top_k=1
    )

    assert indexed["ok"] is True
    assert len(chunks) >= 3
    assert all(len(item["chunk_text"]) <= 900 for item in chunks)
    assert result["data"]["status"] == "found"
    assert len(result["data"]["evidence"]) == 1
    excerpt = result["data"]["evidence"][0]["text_excerpt"]
    assert "连续请假超过三天需要导师签字" in excerpt
    assert "甲" * 50 not in excerpt
    assert "乙" * 50 not in excerpt
    relevant = next(item for item in chunks if "连续请假" in item["chunk_text"])
    assert relevant["metadata"]["paragraph_start"]
    assert relevant["metadata"]["section"] == "第二章 请假规则"


def test_r05_missing_rule_is_not_fabricated_and_arguments_are_bounded(
    document_data_dir: Path,
) -> None:
    file_id, path = _register(document_data_dir, "有限规则.docx", "word", b"")
    document = Document()
    document.add_paragraph("本材料只说明课程选修流程。")
    document.save(path)
    assert index_document(file_id)["ok"] is True

    missing = retrieve_document({"file_ids": [file_id]}, "火星交换生补贴规则")
    arguments, rejected = agent._execute_tool(
        "retrieve_document",
        '{"scope":{"file_ids":["' + file_id + '"],"path":"../data"},"query":"规则"}',
    )

    assert missing["ok"] is True
    assert missing["data"]["status"] == "not_found"
    assert missing["data"]["evidence"] == []
    assert missing["evidence"] == []
    assert missing["message"] == "当前材料中未找到足够依据。"
    assert missing["data"]["result_summary"] == "当前材料中未找到足够依据。"
    assert arguments["scope"]["path"] == "../data"
    assert rejected["error_code"] == "INVALID_TOOL_ARGUMENTS"


def test_deleted_word_cleans_chunks_and_failed_delete_restores_them(
    document_data_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    success_id, success_path = _register(document_data_dir, "删除索引.docx", "word", b"")
    document = Document()
    document.add_paragraph("等待删除的索引文本。")
    document.save(success_path)
    assert index_document(success_id)["ok"] is True

    deleted = delete_uploaded_file(
        {"action_id": "delete-ok", "content": success_id, "target_file": "删除索引.docx"}
    )
    assert deleted["ok"] is True
    assert database.get_document_chunks(success_id) == []
    assert database.get_file_record_by_id(success_id)["index_status"] == "not_required"

    failed_id, failed_path = _register(document_data_dir, "恢复索引.docx", "word", b"")
    document = Document()
    document.add_paragraph("删除失败后必须保留的索引文本。")
    document.save(failed_path)
    assert index_document(failed_id)["ok"] is True
    before = database.get_document_chunks(failed_id)
    original_unlink = Path.unlink

    def reject_staging(path: Path, *args, **kwargs):
        if path.name.startswith(".delete-"):
            raise OSError("simulated")
        return original_unlink(path, *args, **kwargs)

    monkeypatch.setattr(Path, "unlink", reject_staging)
    failed = delete_uploaded_file(
        {"action_id": "delete-fail", "content": failed_id, "target_file": "恢复索引.docx"}
    )

    assert failed["ok"] is False
    assert failed_path.exists()
    assert database.get_document_chunks(failed_id) == before
    restored = database.get_file_record_by_id(failed_id)
    assert restored["lifecycle_status"] == "cleanup_failed"
    assert restored["index_status"] == "indexed"
    restored_evidence = retrieve_document(
        {"file_id": failed_id}, "删除失败后必须保留的索引文本"
    )
    # cleanup_failed files intentionally remain non-queryable until lifecycle repair,
    # while their derived chunks stay synchronized for a later safe resume.
    assert restored_evidence["data"]["status"] == "not_found"


def test_f09_reindex_activation_failure_keeps_old_chunks_and_queryability(
    document_data_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    file_id, path = _register(document_data_dir, "原子重索引失败.docx", "word", b"")
    document = Document()
    document.add_paragraph("旧索引中必须持续可查询的奖学金规则。")
    document.save(path)
    assert index_document(file_id)["ok"] is True
    old_chunks = database.get_document_chunks(file_id)

    replacement = Document()
    replacement.add_paragraph("候选索引的新规则不应在失败后生效。")
    replacement.add_paragraph("候选索引第二段用于触发部分写入。")
    replacement.save(path)

    original_insert = database.DOCUMENT_REPOSITORY._insert

    def fail_after_one_candidate(connection, candidate_file_id, chunks):
        original_insert(connection, candidate_file_id, chunks[:1])
        raise sqlite3.OperationalError("simulated activation failure")

    monkeypatch.setattr(database.DOCUMENT_REPOSITORY, "_insert", fail_after_one_candidate)
    result = reindex_document(file_id)

    assert result["ok"] is False
    assert result["error_code"] == "INDEX_WRITE_ERROR"
    assert result["data"]["active_index_preserved"] is True
    assert database.get_document_chunks(file_id) == old_chunks
    old_result = retrieve_document(
        {"file_id": file_id}, "旧索引中必须持续可查询的奖学金规则"
    )
    assert old_result["data"]["status"] == "found"
    assert database.search_document_chunks(
        file_ids=[file_id], query="候选索引的新规则不应在失败后生效", limit=5
    ) == []
    lifecycle = get_file_lifecycle(file_id)["data"]
    assert lifecycle["status"] == "QUERYABLE"
    assert lifecycle["queryable"] is True
    assert lifecycle["index_status"] == "indexed"
    assert lifecycle["error"]["error_code"] == "INDEX_WRITE_ERROR"
    assert lifecycle["error"]["failure_stage"] == "index"


def test_f09_successful_reindex_atomically_switches_without_mixing_or_duplicates(
    document_data_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    file_id, path = _register(document_data_dir, "原子重索引成功.docx", "word", b"")
    document = Document()
    document.add_paragraph("只属于旧版本的材料内容。")
    document.save(path)
    assert index_document(file_id)["ok"] is True
    old_chunk_ids = {row["chunk_id"] for row in database.get_document_chunks(file_id)}

    replacement = Document()
    replacement.add_paragraph("只属于新版本的材料内容。")
    replacement.add_paragraph("新版本的第二条有效材料。")
    replacement.save(path)

    original_builder = document_index._word_chunks
    observed_during_build: dict[str, object] = {}

    def inspect_active_generation(candidate_file_id: str, candidate_path: Path):
        current = database.get_file_record_by_id(candidate_file_id)
        observed_during_build["status"] = current["status"]
        observed_during_build["old_retrieval"] = retrieve_document(
            {"file_id": candidate_file_id}, "只属于旧版本的材料内容"
        )["data"]["status"]
        return original_builder(candidate_file_id, candidate_path)

    monkeypatch.setattr(document_index, "_word_chunks", inspect_active_generation)
    switched = reindex_document(file_id)
    new_chunks = database.get_document_chunks(file_id)

    assert switched["ok"] is True
    assert observed_during_build == {
        "status": "REINDEXING",
        "old_retrieval": "found",
    }
    assert old_chunk_ids.isdisjoint({row["chunk_id"] for row in new_chunks})
    assert all("旧版本" not in row["chunk_text"] for row in new_chunks)
    assert any("新版本" in row["chunk_text"] for row in new_chunks)
    assert database.search_document_chunks(
        file_ids=[file_id], query="只属于旧版本的材料内容", limit=5
    ) == []
    assert database.search_document_chunks(
        file_ids=[file_id], query="只属于新版本的材料内容", limit=5
    )

    repeated = reindex_document(file_id)
    repeated_chunks = database.get_document_chunks(file_id)
    assert repeated["ok"] is True
    assert repeated_chunks == new_chunks
    assert len({row["chunk_id"] for row in repeated_chunks}) == len(repeated_chunks)


def test_f10_reprocess_parse_failure_preserves_one_complete_active_generation(
    document_data_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    file_id, path = _register(document_data_dir, "安全重新解析.docx", "word", b"")
    document = Document()
    document.add_paragraph("重新解析失败后仍需保留的原始结果。")
    document.save(path)
    assert index_document(file_id)["ok"] is True
    old_chunks = database.get_document_chunks(file_id)

    def fail_parse(file_id: str, path: Path):
        current = database.get_file_record_by_id(file_id)
        assert current["status"] == "REPROCESSING"
        assert retrieve_document(
            {"file_id": file_id}, "重新解析失败后仍需保留的原始结果"
        )["data"]["status"] == "found"
        raise ValueError("simulated parse failure")

    monkeypatch.setattr(document_index, "_word_chunks", fail_parse)
    result = reprocess_document(file_id)

    assert result["ok"] is False
    assert result["error_code"] == "WORD_PARSE_ERROR"
    assert result["data"]["active_index_preserved"] is True
    assert database.get_document_chunks(file_id) == old_chunks
    lifecycle = get_file_lifecycle(file_id)["data"]
    assert lifecycle["status"] == "QUERYABLE"
    assert lifecycle["error"]["failure_stage"] == "parse"
    assert retrieve_document(
        {"file_id": file_id}, "重新解析失败后仍需保留的原始结果"
    )["data"]["status"] == "found"


def test_f10_successful_reprocess_activates_only_the_new_parsed_generation(
    document_data_dir: Path
) -> None:
    file_id, path = _register(document_data_dir, "安全重新解析成功.docx", "word", b"")
    document = Document()
    document.add_paragraph("重新解析前的旧解析结果。")
    document.save(path)
    assert index_document(file_id)["ok"] is True

    replacement = Document()
    replacement.add_paragraph("重新解析后的新解析结果。")
    replacement.save(path)
    result = reprocess_document(file_id)
    chunks = database.get_document_chunks(file_id)

    assert result["ok"] is True
    assert [row["chunk_text"] for row in chunks] == ["重新解析后的新解析结果。"]
    assert database.search_document_chunks(
        file_ids=[file_id], query="重新解析前的旧解析结果", limit=5
    ) == []
    assert database.search_document_chunks(
        file_ids=[file_id], query="重新解析后的新解析结果", limit=5
    )
    lifecycle = get_file_lifecycle(file_id)["data"]
    assert lifecycle["status"] == "QUERYABLE"
    assert lifecycle["error"] is None
