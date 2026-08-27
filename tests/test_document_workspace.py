"""V3 Document Workspace API and deterministic view tests."""

from io import BytesIO
from pathlib import Path

import httpx
import pytest
from docx import Document
from openpyxl import Workbook

from backend import database
from backend.services import document_index
from backend.services.file_upload import save_uploaded_file
from backend.main import app
from backend.tools import excel_utils


pytestmark = pytest.mark.anyio


@pytest.fixture
def anyio_backend() -> str:
    return "asyncio"


@pytest.fixture
async def client():
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as api_client:
        yield api_client


@pytest.fixture
def workspace_data_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)
    return tmp_path


def _word_bytes(text: str) -> bytes:
    output = BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(output)
    return output.getvalue()


def _excel_bytes() -> bytes:
    output = BytesIO()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "学生"
    worksheet.append(["学号", "姓名"])
    worksheet.append(["S001", "张三"])
    workbook.save(output)
    workbook.close()
    return output.getvalue()


async def test_workspace_lists_file_cards_with_required_metadata(
    client: httpx.AsyncClient,
) -> None:
    response = await client.get("/api/files")

    assert response.status_code == 200
    items = response.json()["data"]
    assert items
    assert {
        "file_id",
        "file_name",
        "file_type",
        "size",
        "created_time",
        "status",
        "lifecycle_status",
        "parse_status",
        "index_status",
    } <= set(items[0])


async def test_workspace_searches_by_partial_file_name(
    client: httpx.AsyncClient,
) -> None:
    response = await client.get("/api/files", params={"search": "学生成绩"})

    assert response.status_code == 200
    names = [item["file_name"] for item in response.json()["data"]]
    assert names == ["学生成绩.xlsx"]


async def test_workspace_filters_and_sorts_in_python(
    client: httpx.AsyncClient,
) -> None:
    response = await client.get(
        "/api/files",
        params={
            "file_type": "excel",
            "lifecycle_status": "ready",
            "sort_by": "size",
            "sort_order": "asc",
        },
    )

    assert response.status_code == 200
    items = response.json()["data"]
    assert items
    assert all(item["file_type"] == "excel" for item in items)
    assert all(item["lifecycle_status"] == "ready" for item in items)
    assert [item["size"] for item in items] == sorted(item["size"] for item in items)


async def test_workspace_reads_file_details_by_stable_id(
    client: httpx.AsyncClient,
) -> None:
    record = database.get_file_record("综合评价.docx")

    response = await client.get(f"/api/files/{record['file_id']}")

    assert response.status_code == 200
    detail = response.json()["data"]
    assert detail["file_id"] == record["file_id"]
    assert detail["file_name"] == "综合评价.docx"
    assert detail["lifecycle_status"] == "ready"
    assert detail["parse_status"] == "not_required"
    assert detail["index_status"] == "not_required"


async def test_workspace_quick_preview_is_bounded_and_uses_real_content(
    client: httpx.AsyncClient,
    workspace_data_dir: Path,
) -> None:
    word = save_uploaded_file("预览材料.docx", _word_bytes("用于快速预览的真实 Word 内容。"))
    excel = save_uploaded_file("预览表格.xlsx", _excel_bytes())

    word_response = await client.get(
        f"/api/files/{word['data']['file_id']}/preview"
    )
    excel_response = await client.get(
        f"/api/files/{excel['data']['file_id']}/preview"
    )

    assert word_response.status_code == 200
    word_preview = word_response.json()["data"]
    assert word_preview["preview_type"] == "text"
    assert any("真实 Word 内容" in item["text"] for item in word_preview["items"])
    assert len(word_preview["items"]) <= 10
    assert excel_response.status_code == 200
    excel_preview = excel_response.json()["data"]
    assert excel_preview["preview_type"] == "table"
    assert excel_preview["sheets"] == [
        {
            "sheet_name": "学生",
            "rows": [["学号", "姓名"], ["S001", "张三"]],
        }
    ]


async def test_f09_f10_workspace_calls_safe_reprocess_and_reindex_apis(
    client: httpx.AsyncClient,
    workspace_data_dir: Path,
) -> None:
    uploaded = save_uploaded_file("Workspace重处理.docx", _word_bytes("旧工作区内容。"))
    file_id = uploaded["data"]["file_id"]
    path = workspace_data_dir / "uploads" / "Workspace重处理.docx"

    path.write_bytes(_word_bytes("重新解析后的内容。"))
    reprocessed = await client.post(f"/api/files/{file_id}/reprocess")
    assert reprocessed.status_code == 200
    assert [row["chunk_text"] for row in database.get_document_chunks(file_id)] == [
        "重新解析后的内容。"
    ]

    path.write_bytes(_word_bytes("重新索引后的内容。"))
    reindexed = await client.post(f"/api/files/{file_id}/reindex")
    detail = await client.get(f"/api/files/{file_id}")

    assert reindexed.status_code == 200
    assert [row["chunk_text"] for row in database.get_document_chunks(file_id)] == [
        "重新索引后的内容。"
    ]
    assert detail.json()["data"]["canonical_status"] == "QUERYABLE"
    assert detail.json()["data"]["error_summary"] is None


async def test_workspace_reindex_failure_returns_persisted_error_summary(
    client: httpx.AsyncClient,
    workspace_data_dir: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    uploaded = save_uploaded_file("Workspace失败状态.docx", _word_bytes("旧有效内容。"))
    file_id = uploaded["data"]["file_id"]
    old_chunks = database.get_document_chunks(file_id)

    monkeypatch.setattr(
        document_index,
        "_word_chunks",
        lambda file_id, path: (_ for _ in ()).throw(ValueError("simulated")),
    )
    response = await client.post(f"/api/files/{file_id}/reindex")
    detail = await client.get(f"/api/files/{file_id}")

    assert response.status_code == 500
    assert database.get_document_chunks(file_id) == old_chunks
    assert detail.status_code == 200
    payload = detail.json()["data"]
    assert payload["canonical_status"] == "QUERYABLE"
    assert payload["error_summary"]["error_code"] == "WORD_PARSE_ERROR"
    assert payload["error_summary"]["failure_stage"] == "parse"
