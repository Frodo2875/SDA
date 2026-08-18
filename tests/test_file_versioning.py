"""V2.9 W01-W07 Diff, Version, Undo, Rollback, and HITL safety tests."""

import sqlite3
import shutil
from pathlib import Path

import pytest
from docx import Document

from backend import database
from backend.services import confirmation
from backend.services.confirmation import (
    cancel_action,
    confirm_action,
    create_pending_action,
    create_pending_undo_action,
    rollback_file,
)
from backend.services.file_versioning import preview_word_diff
from backend.tools import excel_utils


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CONTENT_ONE = "V2.9 第一条版本化综合评价。"
CONTENT_TWO = "V2.9 第二条版本化综合评价。"


@pytest.fixture
def versioned_word(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> tuple[str, Path]:
    target = tmp_path / "综合评价.docx"
    shutil.copy2(PROJECT_ROOT / "data" / "综合评价.docx", target)
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)
    record = database.get_file_record("综合评价.docx")
    assert record is not None
    return record["file_id"], target


def _paragraphs(path: Path) -> list[str]:
    return [item.text for item in Document(path).paragraphs if item.text]


def _pending_append(content: str, session_id: str = "version-test") -> dict:
    result = create_pending_action(
        session_id=session_id,
        target_file="综合评价.docx",
        student_id="S001",
        student_name="张三",
        content=content,
    )
    assert result["ok"] is True
    return result["data"]


def _confirmed_append(content: str, session_id: str = "version-test") -> dict:
    action = _pending_append(content, session_id)
    result = confirm_action(action["action_id"])
    assert result["ok"] is True
    return result


def test_w01_diff_preview_has_no_file_or_version_side_effect(versioned_word) -> None:
    file_id, target = versioned_word
    before = target.read_bytes()

    result = preview_word_diff(
        file_id, {"operation_type": "append", "content": CONTENT_ONE}
    )

    assert result["ok"] is True
    assert result["data"]["target_file"] == "综合评价.docx"
    assert result["data"]["target_object"] == "Word 文档正文"
    assert result["data"]["location"] == "文档末尾"
    assert result["data"]["operation_type"] == "append"
    assert result["data"]["after"].endswith(CONTENT_ONE)
    assert result["data"]["impact_scope"] == "新增 1 个段落"
    assert target.read_bytes() == before
    assert database.list_file_versions(file_id) == []


def test_w02_pending_write_freezes_diff_but_changes_nothing(versioned_word) -> None:
    file_id, target = versioned_word
    before = target.read_bytes()

    action = _pending_append(CONTENT_ONE, "w02")

    assert action["status"] == "pending"
    assert action["file_id"] == file_id
    assert action["operation"]["operation_type"] == "append"
    assert action["diff_preview"]["after"].endswith(CONTENT_ONE)
    assert target.read_bytes() == before
    assert database.list_file_versions(file_id) == []


def test_w03_confirm_creates_parent_and_new_version_once(versioned_word) -> None:
    file_id, target = versioned_word
    action = _pending_append(CONTENT_ONE, "w03")

    first = confirm_action(action["action_id"])
    versions = database.list_file_versions(file_id)
    bytes_after_first = target.read_bytes()
    second = confirm_action(action["action_id"])

    assert first["ok"] is True
    assert _paragraphs(target).count(CONTENT_ONE) == 1
    assert [item["version_number"] for item in versions] == [1, 2]
    assert versions[0]["change_type"] == "snapshot"
    assert versions[1]["change_type"] == "append"
    assert versions[1]["parent_version_id"] == versions[0]["version_id"]
    assert all(Path(database.DB_PATH.parent / item["storage_path"].removeprefix("data/")).is_file() for item in versions)
    assert second["ok"] is False
    assert second["error_code"] == "ACTION_NOT_PENDING"
    assert target.read_bytes() == bytes_after_first
    assert len(database.list_file_versions(file_id)) == 2


def test_w04_cancel_write_keeps_file_and_creates_no_version(versioned_word) -> None:
    file_id, target = versioned_word
    before = target.read_bytes()
    action = _pending_append(CONTENT_ONE, "w04")

    result = cancel_action(action["action_id"])

    assert result["ok"] is True
    assert target.read_bytes() == before
    assert database.list_file_versions(file_id) == []


def test_w05_undo_requires_confirmation_and_creates_new_history(versioned_word) -> None:
    file_id, target = versioned_word
    original = target.read_bytes()
    _confirmed_append(CONTENT_ONE, "w05-write")
    modified = target.read_bytes()
    versions_before = database.list_file_versions(file_id)

    pending = create_pending_undo_action(session_id="w05-undo", file_id=file_id)
    assert pending["ok"] is True
    assert pending["data"]["diff_preview"]["operation_type"] == "undo"
    assert target.read_bytes() == modified
    confirmed = confirm_action(pending["data"]["action_id"])
    versions_after = database.list_file_versions(file_id)

    assert confirmed["ok"] is True
    assert target.read_bytes() == original
    assert len(versions_before) == 2
    assert len(versions_after) == 3
    assert versions_after[-1]["change_type"] == "undo"
    assert versions_after[-1]["parent_version_id"] == versions_before[-1]["version_id"]
    assert database.get_task_record(pending["data"]["task_id"])["status"] == "success"


def test_w06_rollback_cancel_then_confirm_is_safe(versioned_word) -> None:
    file_id, target = versioned_word
    original = target.read_bytes()
    _confirmed_append(CONTENT_ONE, "w06-one")
    _confirmed_append(CONTENT_TWO, "w06-two")
    modified = target.read_bytes()
    versions = database.list_file_versions(file_id)
    baseline_id = versions[0]["version_id"]

    cancelled = rollback_file(file_id, baseline_id, session_id="w06-cancel")
    assert cancelled["ok"] is True
    assert target.read_bytes() == modified
    assert cancel_action(cancelled["data"]["action_id"])["ok"] is True
    assert target.read_bytes() == modified
    assert len(database.list_file_versions(file_id)) == len(versions)

    pending = rollback_file(file_id, baseline_id, session_id="w06-confirm")
    assert pending["data"]["target_version_id"] == baseline_id
    result = confirm_action(pending["data"]["action_id"])
    current = database.list_file_versions(file_id)

    assert result["ok"] is True
    assert target.read_bytes() == original
    assert len(current) == len(versions) + 1
    assert current[-1]["change_type"] == "rollback"


def test_w07_write_failure_restores_file_and_leaves_no_half_version(
    versioned_word, monkeypatch: pytest.MonkeyPatch
) -> None:
    file_id, target = versioned_word
    original = target.read_bytes()
    action = _pending_append(CONTENT_ONE, "w07")

    def fail_commit(**kwargs):
        raise sqlite3.OperationalError("forced version commit failure")

    monkeypatch.setattr(database, "commit_file_version_transition", fail_commit)
    result = confirm_action(action["action_id"])

    assert result["ok"] is False
    assert target.read_bytes() == original
    assert database.list_file_versions(file_id) == []
    version_files = list((database.DB_PATH.parent / "versions").rglob("*.docx"))
    assert version_files == []


def test_w07_rollback_failure_keeps_current_file_and_history(
    versioned_word, monkeypatch: pytest.MonkeyPatch
) -> None:
    file_id, target = versioned_word
    _confirmed_append(CONTENT_ONE, "w07-rollback-seed")
    current_bytes = target.read_bytes()
    versions = database.list_file_versions(file_id)
    pending = rollback_file(file_id, versions[0]["version_id"], session_id="w07-rollback")
    from backend.services import file_versioning

    monkeypatch.setattr(
        file_versioning,
        "_atomic_replace_bytes",
        lambda path, content: (_ for _ in ()).throw(OSError("forced rollback failure")),
    )
    result = confirm_action(pending["data"]["action_id"])

    assert result["ok"] is False
    assert target.read_bytes() == current_bytes
    assert database.list_file_versions(file_id) == versions


def test_w07_read_only_system_file_cannot_be_changed(
    versioned_word, isolated_database: Path
) -> None:
    file_id, target = versioned_word
    before = target.read_bytes()
    with sqlite3.connect(isolated_database) as connection:
        connection.execute("UPDATE files SET writable = 0 WHERE file_id = ?", (file_id,))
        connection.commit()

    result = preview_word_diff(
        file_id, {"operation_type": "append", "content": CONTENT_ONE}
    )

    assert result["ok"] is False
    assert result["error_code"] == "FILE_WRITE_FORBIDDEN"
    assert target.read_bytes() == before
    assert database.list_file_versions(file_id) == []


def test_w07_confirm_rejects_file_changed_after_preview(versioned_word) -> None:
    file_id, target = versioned_word
    action = _pending_append(CONTENT_ONE, "w07-stale-diff")
    document = Document(target)
    document.add_paragraph("预览后由外部流程产生的变化。")
    document.save(target)
    externally_changed = target.read_bytes()

    result = confirm_action(action["action_id"])

    assert result["ok"] is False
    assert result["data"]["action_result"]["error_code"] == "FILE_CHANGED_SINCE_PREVIEW"
    assert target.read_bytes() == externally_changed
    assert database.list_file_versions(file_id) == []
