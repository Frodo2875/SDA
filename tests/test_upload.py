"""Integration tests for safe Excel and Word uploads."""

from io import BytesIO
from pathlib import Path

import httpx
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from backend import database
from backend.main import app
from backend.services.file_upload import save_uploaded_file
from backend.tools import excel_utils
from backend.tools.file_tools import list_files


pytestmark = pytest.mark.anyio


@pytest.fixture
def anyio_backend() -> str:
    return "asyncio"


@pytest.fixture
async def client():
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as api_client:
        yield api_client


@pytest.fixture
def upload_data_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)
    return tmp_path


def _xlsx_bytes() -> bytes:
    output = BytesIO()
    workbook = Workbook()
    workbook.active.append(["测试字段"])
    workbook.active.append(["测试值"])
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _docx_bytes() -> bytes:
    output = BytesIO()
    document = Document()
    document.add_paragraph("测试上传文档")
    document.save(output)
    return output.getvalue()


async def test_upload_valid_excel_saves_registers_and_lists_file(
    client: httpx.AsyncClient, upload_data_dir: Path
) -> None:
    response = await client.post(
        "/api/files/upload",
        files={
            "file": (
                "补充材料.xlsx",
                _xlsx_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["ok"] is True
    assert body["data"]["path"] == "data/uploads/补充材料.xlsx"
    saved_path = upload_data_dir / "uploads" / "补充材料.xlsx"
    assert saved_path.is_file()
    workbook = load_workbook(saved_path, read_only=True)
    workbook.close()

    record = database.get_file_record("补充材料.xlsx")
    assert record["file_type"] == "excel"
    assert record["file_path"] == "data/uploads/补充材料.xlsx"
    listed = list_files()
    assert [(item["file_name"], item["path"]) for item in listed["data"]] == [
        ("补充材料.xlsx", "data/uploads/补充材料.xlsx")
    ]


async def test_upload_valid_word_can_be_opened(
    client: httpx.AsyncClient, upload_data_dir: Path
) -> None:
    response = await client.post(
        "/api/files/upload",
        files={"file": ("获奖材料.docx", _docx_bytes(), "application/octet-stream")},
    )

    assert response.status_code == 200
    saved_path = upload_data_dir / "uploads" / "获奖材料.docx"
    assert Document(saved_path).paragraphs[0].text == "测试上传文档"
    record = database.get_file_record("获奖材料.docx")
    assert record["file_type"] == "word"
    assert record["writable"] == 1


async def test_pdf_is_rejected_without_creating_file(
    client: httpx.AsyncClient, upload_data_dir: Path
) -> None:
    response = await client.post(
        "/api/files/upload",
        files={"file": ("材料.pdf", b"%PDF-1.4 test", "application/pdf")},
    )

    assert response.status_code == 400
    assert response.json()["error_code"] == "UNSUPPORTED_FILE_TYPE"
    assert "暂不支持 PDF" in response.json()["message"]
    assert not (upload_data_dir / "uploads" / "材料.pdf").exists()


@pytest.mark.parametrize("file_name", ["../逃逸.xlsx", "..\\逃逸.xlsx", "folder/逃逸.docx"])
def test_path_traversal_file_names_are_rejected(
    file_name: str, upload_data_dir: Path
) -> None:
    result = save_uploaded_file(file_name, _xlsx_bytes())

    assert result["ok"] is False
    assert result["error_code"] == "INVALID_FILE_NAME"
    assert list(upload_data_dir.rglob("逃逸*")) == []


async def test_duplicate_name_is_rejected_without_overwrite(
    client: httpx.AsyncClient, upload_data_dir: Path
) -> None:
    first_content = _xlsx_bytes()
    first = await client.post(
        "/api/files/upload", files={"file": ("重复.xlsx", first_content)}
    )
    target = upload_data_dir / "uploads" / "重复.xlsx"
    before = target.read_bytes()
    second = await client.post(
        "/api/files/upload", files={"file": ("重复.xlsx", _xlsx_bytes())}
    )

    assert first.status_code == 200
    assert second.status_code == 409
    assert second.json()["error_code"] == "FILE_ALREADY_EXISTS"
    assert target.read_bytes() == before


@pytest.mark.parametrize(
    ("file_name", "content"),
    [("损坏.xlsx", b"not an excel file"), ("损坏.docx", b"not a word file")],
)
async def test_corrupt_office_file_is_rejected_and_not_registered(
    client: httpx.AsyncClient,
    upload_data_dir: Path,
    file_name: str,
    content: bytes,
) -> None:
    response = await client.post(
        "/api/files/upload", files={"file": (file_name, content)}
    )

    assert response.status_code == 400
    assert response.json()["error_code"] == "INVALID_FILE_CONTENT"
    assert not (upload_data_dir / "uploads" / file_name).exists()
    assert database.get_file_record(file_name) is None


async def test_openapi_exposes_upload_route(client: httpx.AsyncClient) -> None:
    paths = (await client.get("/openapi.json")).json()["paths"]

    assert "/api/files/upload" in paths
