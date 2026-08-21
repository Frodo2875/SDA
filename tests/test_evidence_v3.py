"""V3.7 Evidence 2.0 source-location tests for PDF, Excel and Word."""

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from backend import database
from backend.document_blocks import block_to_chunks, make_document_block
from backend.evidence import is_evidence_locatable
from backend.services.document_index import index_document, retrieve_document
from backend.tools import excel_utils
from backend.tools.schema_tools import inspect_excel
from backend.tools.table_tools import query_table


def _register(
    directory: Path,
    *,
    name: str,
    file_type: str,
    content: bytes = b"fixture",
) -> tuple[str, Path]:
    upload_dir = directory / "uploads"
    upload_dir.mkdir(exist_ok=True)
    path = upload_dir / name
    path.write_bytes(content)
    database.register_file(
        file_name=name,
        file_type=file_type,
        file_path=f"data/uploads/{name}",
        lifecycle_status="ready",
        parse_status="parsed",
        queryable=True,
        index_status="indexed" if file_type in {"pdf", "word"} else "not_required",
    )
    return database.get_file_record(name)["file_id"], path


def test_pdf_evidence_locates_page_block_bbox_and_confidence(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)
    file_id, _ = _register(tmp_path, name="扫描证明.pdf", file_type="pdf")
    block = make_document_block(
        file_id=file_id,
        sequence=0,
        block_type="paragraph",
        content="原文明确规定申请人需要提交成绩证明。",
        page_no=3,
        bbox=[12.0, 24.0, 220.0, 58.0],
        confidence=0.93,
    )
    database.replace_document_chunks(
        file_id, block_to_chunks(block, source_type="ocr")
    )

    result = retrieve_document({"file_id": file_id}, "提交成绩证明", top_k=1)

    evidence = result["data"]["evidence"][0]
    assert is_evidence_locatable(evidence) is True
    assert evidence["file_id"] == file_id
    assert evidence["page_no"] == 3
    assert evidence["block_id"] == block.block_id
    assert evidence["bbox"] == [12.0, 24.0, 220.0, 58.0]
    assert evidence["confidence"] == 0.93
    chunk = database.get_document_chunks(file_id)[0]
    assert chunk["metadata"]["block"]["content"] == block.content


def test_excel_evidence_locates_table_and_exact_cell(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)
    path = tmp_path / "学生成绩.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "成绩表"
    sheet.append(["学号", "平均成绩"])
    sheet.append(["S001", 92])
    workbook.save(path)
    workbook.close()
    database.initialize_database()
    record = database.get_file_record(path.name)
    assert inspect_excel(record["file_id"])["ok"] is True

    result = query_table(
        record["file_id"],
        "成绩表",
        filters=[{"field": "学号", "op": "=", "value": "S001"}],
        select=["平均成绩"],
    )

    evidence = result["evidence_chain"][0]
    assert is_evidence_locatable(evidence) is True
    assert evidence["file_id"] == record["file_id"]
    assert evidence["table"] == "成绩表"
    assert evidence["cell"] == "B2"
    assert evidence["sheet"] == "成绩表"
    assert evidence["field"] == "平均成绩"
    assert 0.0 <= evidence["confidence"] <= 1.0
    source = load_workbook(path, read_only=True, data_only=True)
    try:
        assert str(source["成绩表"][evidence["cell"]].value) == evidence["value_summary"]
    finally:
        source.close()


def test_word_evidence_locates_original_block(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)
    file_id, path = _register(
        tmp_path, name="申请指南.docx", file_type="word", content=b""
    )
    document = Document()
    document.add_heading("申请要求", level=1)
    document.add_paragraph("申请材料必须由导师签字后提交学院。")
    document.save(path)
    assert index_document(file_id)["ok"] is True

    result = retrieve_document({"file_id": file_id}, "导师签字", top_k=1)

    evidence = result["data"]["evidence"][0]
    assert is_evidence_locatable(evidence) is True
    assert evidence["file_id"] == file_id
    assert evidence["block_id"]
    chunks = database.get_document_chunks(file_id)
    source_chunk = next(
        chunk for chunk in chunks if chunk["metadata"].get("block_id") == evidence["block_id"]
    )
    assert source_chunk["metadata"]["block_type"] == "paragraph"
    assert source_chunk["metadata"]["block"]["content"] == (
        "申请材料必须由导师签字后提交学院。"
    )

