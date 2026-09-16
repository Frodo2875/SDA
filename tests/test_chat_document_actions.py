"""All chat modes expose document proposals without granting model execution."""

import asyncio
from pathlib import Path
from typing import Any

import httpx
import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

from backend import agent, database, main
from backend.schemas import ResearchTaskRequest
from backend.services.research_task import create_research_task, run_research_task
from backend.tools import excel_utils
from frontend import api_client


@pytest.fixture
def source(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> dict[str, Any]:
    monkeypatch.setattr(excel_utils, "DATA_DIR", tmp_path)
    database.save_chat_message(session_id="draft", role="assistant", content="已核实的分析正文",
                               used_tools=[], status="completed")
    return {"session_id": "draft", "source_answer": "已核实的分析正文", "content": "用户编辑后的正文"}


def request(method: str, path: str, **kwargs: Any) -> httpx.Response:
    async def call() -> httpx.Response:
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://test") as client:
            return await client.request(method, path, **kwargs)
    return asyncio.run(call())


@pytest.mark.parametrize("format", ["md", "docx"])
@pytest.mark.parametrize("decision", ["confirm", "cancel"])
def test_new_document_requires_real_confirmation(source, tmp_path, format, decision) -> None:
    response = request("POST", "/api/document-actions/preview", json={**source, "format": format})
    assert response.status_code == 200, response.text
    action = response.json()["data"]
    path = tmp_path / "uploads" / action["target_file"]
    assert action["status"] == "pending"
    assert not path.exists() and not database.list_file_versions(action["file_id"])
    assert action["diff_preview"]["after"] == source["content"]
    decided = request("POST", f"/api/actions/{action['action_id']}/{decision}")
    assert decided.status_code == 200, decided.text
    if decision == "confirm":
        text = path.read_text() if format == "md" else "\n".join(p.text for p in Document(path).paragraphs)
        assert source["content"] in text
        assert len(database.list_file_versions(action["file_id"])) == 2
    else:
        assert not path.exists() and not database.list_file_versions(action["file_id"])
    assert request("POST", f"/api/actions/{action['action_id']}/confirm").status_code == 409


@pytest.mark.parametrize("extra", [{"confirmed": True}, {"format": "xlsx"}, {"file_path": "../../secret.docx"}, {"content": " "}])
def test_rejects_auto_approval_paths_and_unsupported_formats(source, extra, tmp_path) -> None:
    assert request("POST", "/api/document-actions/preview", json={**source, **extra}).status_code == 422
    assert not list(tmp_path.rglob("report-*"))


@pytest.mark.parametrize("change", [{"session_id": "another"}, {"source_answer": "虚构回答"}])
def test_source_must_belong_to_current_session(source, change) -> None:
    assert request("POST", "/api/document-actions/preview", json={**source, **change}).status_code == 409


def test_visual_guard_uses_server_evidence(source, tmp_path) -> None:
    database.save_chat_message(session_id="draft", role="assistant", content=source["source_answer"], status="completed",
        used_tools=[{"name": "retrieve_document", "result": {"evidence": [{"evidence_id": "unsafe", "locator_type": "image", "evidence_status": "conflict"}]}}])
    response = request("POST", "/api/document-actions/preview", json=source)
    assert response.status_code == 409
    assert response.json()["error_code"] == "VISUAL_FIELD_REVIEW_REQUIRED"
    assert not list(tmp_path.rglob("report-*"))


def test_append_preserves_original_and_blocks_changed_file(source, tmp_path) -> None:
    path = tmp_path / "原文.docx"
    document = Document()
    document.add_paragraph("原始内容")
    document.save(path)
    database.register_file(file_name=path.name, file_type="word", file_path="data/原文.docx", writable=True)
    file_id = database.get_file_record(path.name)["file_id"]
    before = path.read_bytes()
    response = request("POST", "/api/document-actions/preview", json={**source, "file_id": file_id})
    assert response.status_code == 200, response.text
    assert path.read_bytes() == before
    action = response.json()["data"]
    document.add_paragraph("外部修改")
    document.save(path)
    changed = path.read_bytes()
    result = request("POST", f"/api/actions/{action['action_id']}/confirm")
    assert not result.json()["ok"]
    assert path.read_bytes() == changed


@pytest.mark.parametrize("writable", [True, False])
def test_existing_document_append_respects_file_permission(source, tmp_path, writable) -> None:
    path = tmp_path / "资料.docx"
    document = Document()
    document.add_paragraph("保留原文")
    document.save(path)
    database.register_file(file_name=path.name, file_type="word", file_path="data/资料.docx", writable=writable)
    file_id = database.get_file_record(path.name)["file_id"]
    original = path.read_bytes()
    response = request("POST", "/api/document-actions/preview", json={**source, "file_id": file_id})
    assert path.read_bytes() == original
    if not writable:
        assert response.status_code == 409
        assert response.json()["error_code"] == "FILE_WRITE_FORBIDDEN"
        return
    action = response.json()["data"]
    assert request("POST", f"/api/actions/{action['action_id']}/confirm").status_code == 200
    assert [p.text for p in Document(path).paragraphs] == ["保留原文", source["content"]]


def test_research_result_reuses_same_approval(source, monkeypatch) -> None:
    async def answer(*args: Any, **kwargs: Any) -> dict[str, Any]:
        return {"status": "completed", "answer": source["source_answer"], "unified_evidence": []}
    monkeypatch.setattr(agent, "run_agent", answer)
    task = create_research_task(ResearchTaskRequest(session_id="draft", query="研究专业分布并输出报告"))
    payload = {**source, "research_task_id": task["id"]}
    assert request("POST", "/api/document-actions/preview", json=payload).status_code == 409
    run_research_task(task["id"])
    response = request("POST", "/api/document-actions/preview", json=payload)
    assert response.status_code == 200, response.text
    assert response.json()["data"]["status"] == "pending"
    assert request("POST", "/api/document-actions/preview", json={**payload, "session_id": "other"}).status_code == 409


@pytest.mark.parametrize("mode", ["普通问答", "深度研究", "报告生成"])
def test_each_mode_exposes_edit_preview_then_separate_confirmation(monkeypatch, mode) -> None:
    monkeypatch.setattr(api_client, "list_files", lambda **kwargs: [])
    monkeypatch.setattr(api_client, "list_tasks", lambda *args: [])
    monkeypatch.setattr(api_client, "get_traces", lambda **kwargs: [])
    calls = []
    def preview(**payload: Any) -> dict[str, Any]:
        calls.append(payload)
        return {"ok": True, "message": "等待确认", "data": {"action_id": "proposal", "status": "pending", "content": payload["content"], "target_file": "结果.docx"}}
    monkeypatch.setattr(api_client, "preview_document_draft", preview)
    monkeypatch.setattr(api_client, "decide_action", lambda *args: pytest.fail("preview must never confirm"))
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "frontend/app.py").run(timeout=10)
    app.button(key="nav-chat").click().run()
    app.radio(key="chat_mode").set_value(mode).run()
    message = {"role": "assistant", "content": "原回答", "response_status": "completed"}
    if mode != "普通问答":
        message.update(research_task_id="research", research={"status": "COMPLETED", "result": {"status": "completed"}})
    app.session_state.messages = [message]
    app.run()
    next(area for area in app.text_area if area.label == "待保存内容").set_value("编辑后的正文").run()
    app.button(key="document-preview").click().run()
    assert not app.exception
    assert calls[0]["content"] == "编辑后的正文"
    assert calls[0]["source_answer"] == "原回答"
    assert app.button(key="message-1-confirm")
    assert app.button(key="message-1-cancel")
    assert not app.toggle(key="chat-show-sources").value


@pytest.mark.parametrize("operation,method", [("删除文件", "prepare_delete"), ("撤销最近修改", "prepare_undo"), ("恢复历史版本", "prepare_rollback")])
def test_management_controls_only_prepare_actions(monkeypatch, operation, method) -> None:
    file = {"file_id": "file", "file_name": "资料.docx", "writable": True, "file_type": "word",
            "deletable": True, "source_type": "upload", "lifecycle_status": "ready"}
    monkeypatch.setattr(api_client, "list_files", lambda **kwargs: [file])
    monkeypatch.setattr(api_client, "list_tasks", lambda *args: [])
    monkeypatch.setattr(api_client, "list_versions", lambda *args: [{"version_id": "version", "version_number": 1}])
    monkeypatch.setattr(api_client, "decide_action", lambda *args: pytest.fail("must wait for user confirmation"))
    calls = []
    monkeypatch.setattr(api_client, method, lambda *args: calls.append(args) or {"ok": True, "data": {"action_id": "action", "status": "pending"}})
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "frontend/app.py").run(timeout=10)
    app.button(key="nav-chat").click().run()
    app.session_state.files = [file]
    app.selectbox(key="document-operation").set_value(operation).run()
    app.button(key="document-manage-preview").click().run()
    assert not app.exception
    assert calls == [("file", "version", app.session_state.session_id)] if method == "prepare_rollback" else calls == [("file", app.session_state.session_id)]
    assert app.button(key="message-0-confirm")
