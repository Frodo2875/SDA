"""Validation and safe persistence for supported document uploads."""

import os
import shutil
import sqlite3
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook

from backend import database
from backend.services.file_lifecycle import process_uploaded_file
from backend.services.document_index import index_document, route_pdf_file, validate_pdf_file
from backend.services.input_router import (
    route_document_input,
    validate_declared_mime,
    validate_multiformat_content,
)
from backend.tools import excel_utils
from backend.tools.excel_utils import failure, success


SUPPORTED_UPLOADS = {
    ".xlsx": "excel",
    ".docx": "word",
    ".pdf": "pdf",
    ".jpg": "image",
    ".jpeg": "image",
    ".png": "image",
    ".ppt": "presentation",
    ".pptx": "presentation",
    ".md": "markdown",
    ".markdown": "markdown",
    ".txt": "txt",
    ".json": "json",
    ".csv": "csv",
}


def _safe_file_name(file_name: str) -> tuple[str | None, dict[str, Any] | None]:
    if not isinstance(file_name, str) or not file_name.strip():
        return None, failure("INVALID_FILE_NAME", "上传文件名不能为空")
    clean_name = file_name.strip()
    if (
        clean_name in {".", ".."}
        or "/" in clean_name
        or "\\" in clean_name
        or "\x00" in clean_name
        or Path(clean_name).name != clean_name
    ):
        return None, failure("INVALID_FILE_NAME", "文件名不能包含目录路径")
    if Path(clean_name).suffix.lower() not in SUPPORTED_UPLOADS:
        return None, failure(
            "UNSUPPORTED_FILE_TYPE",
            "支持 .md、.markdown、.xlsx、.docx、.pdf、.ppt、.pptx、.txt、.json、.csv、.jpg、.jpeg 和 .png 文件",
        )
    return clean_name, None


def _validate_document(path: Path, suffix: str) -> dict[str, Any] | None:
    try:
        if suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            workbook.close()
        elif suffix == ".docx":
            Document(path)
        elif suffix == ".pdf":
            return validate_pdf_file(path)
        elif suffix in {".ppt", ".pptx", ".txt", ".json", ".csv", ".md", ".markdown"}:
            return validate_multiformat_content(path)
        else:
            routed = route_document_input(path)
            return None if routed["ok"] else routed
    except Exception:
        label = {
            ".xlsx": "Excel",
            ".docx": "Word",
            ".pdf": "PDF",
            ".jpg": "JPEG",
            ".jpeg": "JPEG",
            ".png": "PNG",
            ".ppt": "PPT",
            ".pptx": "PPTX",
            ".md": "Markdown",
            ".markdown": "Markdown",
            ".txt": "TXT",
            ".json": "JSON",
            ".csv": "CSV",
        }[suffix]
        return failure(
            "INVALID_FILE_CONTENT",
            f"文件不是可正常打开的 {label} 文档，请检查文件是否损坏或扩展名是否正确",
        )
    return None


def save_uploaded_file(
    file_name: str,
    content: bytes,
    declared_mime_type: str | None = None,
) -> dict[str, Any]:
    """Validate and save one new upload without ever overwriting a file."""
    clean_name, name_error = _safe_file_name(file_name)
    if name_error is not None:
        return name_error
    if not isinstance(content, bytes) or not content:
        return failure("EMPTY_UPLOAD", "上传文件不能为空")

    data_dir = excel_utils.DATA_DIR
    upload_dir = data_dir / "uploads"
    destination = upload_dir / clean_name
    if (data_dir / clean_name).exists() or destination.exists():
        return failure("FILE_ALREADY_EXISTS", f"同名文件已存在：{clean_name}，禁止覆盖")
    if database.get_file_record(clean_name) is not None:
        return failure("FILE_ALREADY_EXISTS", f"同名文件已登记：{clean_name}，禁止覆盖")

    upload_dir.mkdir(parents=True, exist_ok=True)
    suffix = destination.suffix.lower()
    temporary_path: Path | None = None
    destination_created = False
    route_data: dict[str, Any] | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb",
            dir=upload_dir,
            prefix=".upload-",
            suffix=suffix,
            delete=False,
        ) as temporary_file:
            temporary_file.write(content)
            temporary_file.flush()
            os.fsync(temporary_file.fileno())
            temporary_path = Path(temporary_file.name)

        validation_error = _validate_document(temporary_path, suffix)
        if validation_error is not None:
            return validation_error
        routed = (
            route_pdf_file(temporary_path)
            if suffix == ".pdf"
            else route_document_input(
                temporary_path,
                declared_mime_type=declared_mime_type,
            )
        )
        if suffix == ".pdf" and routed["ok"]:
            mime_error = validate_declared_mime(suffix, declared_mime_type)
            if mime_error is not None:
                routed = mime_error
        if not routed["ok"]:
            return routed
        route_data = routed["data"]

        try:
            descriptor = os.open(
                destination,
                os.O_WRONLY | os.O_CREAT | os.O_EXCL,
                0o600,
            )
        except FileExistsError:
            return failure(
                "FILE_ALREADY_EXISTS", f"同名文件已存在：{clean_name}，禁止覆盖"
            )
        destination_created = True
        with os.fdopen(descriptor, "wb") as output_file, temporary_path.open("rb") as source:
            shutil.copyfileobj(source, output_file)
            output_file.flush()
            os.fsync(output_file.fileno())

        file_type = SUPPORTED_UPLOADS[suffix]
        database.register_file(
            file_name=clean_name,
            file_type=file_type,
            file_path=f"data/uploads/{clean_name}",
            status="active",
            writable=file_type == "word",
            lifecycle_status="uploaded",
            parse_status="pending",
            queryable=False,
            index_status="not_required",
        )
    except sqlite3.IntegrityError:
        if destination_created:
            destination.unlink(missing_ok=True)
        return failure("FILE_ALREADY_EXISTS", f"同名文件已登记：{clean_name}，禁止覆盖")
    except sqlite3.Error:
        if destination_created:
            destination.unlink(missing_ok=True)
        return failure("FILE_REGISTER_ERROR", "文件登记到数据库失败，未保留上传文件")
    except OSError:
        if destination_created:
            destination.unlink(missing_ok=True)
        return failure("FILE_SAVE_ERROR", "上传文件保存失败")
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)

    file_record = database.get_file_record(clean_name)
    if file_record is None:
        return failure("FILE_REGISTER_ERROR", "文件登记结果无法读取")

    if file_type == "image":
        index_result = index_document(file_record["file_id"])
        if not index_result["ok"]:
            return index_result
    else:
        parse_result = process_uploaded_file(
            file_record["file_id"],
            destination,
            _validate_document,
        )
        if not parse_result["ok"]:
            return parse_result

        if file_type in {"word", "pdf", "txt", "json", "presentation", "markdown"}:
            index_result = index_document(file_record["file_id"])
            if not index_result["ok"]:
                return index_result
        elif file_type == "csv":
            from backend.tools.schema_tools import inspect_excel

            schema_result = inspect_excel(file_record["file_id"])
            if not schema_result["ok"]:
                return schema_result

    file_record = database.get_file_record_by_id(file_record["file_id"])
    return success(
        {
            "file_id": file_record["file_id"],
            "file_name": clean_name,
            "file_type": file_type,
            "path": f"data/uploads/{clean_name}",
            "size": len(content),
            "exists": True,
            "status": "active",
            "source_type": file_record["source_type"],
            "trust_level": "unknown",
            "lifecycle_status": file_record["lifecycle_status"],
            "parse_status": file_record["parse_status"],
            "queryable": bool(file_record["queryable"]),
            "index_status": file_record["index_status"],
            "input_route": (route_data or {}).get("route"),
            "mime_type": (route_data or {}).get("mime_type"),
            "image_format": (route_data or {}).get("image_format"),
        },
        "文件上传并校验成功",
    )
