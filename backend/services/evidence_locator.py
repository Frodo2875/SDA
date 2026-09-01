"""Resolve persisted Evidence 2.0 IDs into bounded, read-only previews."""

import re
import time
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

from backend import database
from backend.services.file_locator import FileLocatorError, resolve_by_file_id
from backend.services.input_router import build_image_preview
from backend.services.multiformat_parser import parse_csv_table
from backend.services.trace_service import record_trace
from backend.tools.excel_utils import failure, success


EVIDENCE_ID_PATTERN = re.compile(r"^[A-Za-z0-9_-]{1,128}$")


def locate_evidence(
    evidence_id: str, *, session_id: str | None = None
) -> dict[str, Any]:
    """Return a trusted locator and preview for one actual answer Evidence."""
    clean_id = str(evidence_id or "").strip()
    trace_session = str(session_id or "evidence-preview")[:128]
    started = time.perf_counter()
    if not EVIDENCE_ID_PATTERN.fullmatch(clean_id):
        return _location_failure(
            clean_id, trace_session, "INVALID_EVIDENCE_ID", "Evidence 标识无效"
        )
    evidence = database.get_evidence_location(clean_id)
    if evidence is None:
        return _location_failure(
            clean_id,
            trace_session,
            "EVIDENCE_LOCATION_NOT_FOUND",
            "来源存在，但当前无法打开对应预览位置",
        )
    record = database.get_file_record_by_id(str(evidence.get("file_id") or ""))
    if record is None:
        return _location_failure(
            clean_id,
            trace_session,
            "EVIDENCE_SOURCE_UNAVAILABLE",
            "来源存在，但当前无法打开对应预览位置",
        )
    try:
        path = resolve_by_file_id(record["file_id"])
        if evidence.get("locator_type") == "visual_table":
            location = _locate_visual_table(evidence, record, path)
        elif record["file_type"] == "excel":
            location = _locate_excel(evidence, record, path)
        elif record["file_type"] == "csv":
            location = _locate_csv(evidence, record, path)
        elif record["file_type"] in {"pdf", "word", "presentation", "txt", "json"}:
            location = _locate_document(evidence, record)
        elif record["file_type"] == "image":
            location = _locate_image(evidence, record, path)
        else:
            raise ValueError("不支持的 Evidence 文件类型")
    except (FileLocatorError, KeyError, TypeError, ValueError, OSError):
        return _location_failure(
            clean_id,
            trace_session,
            "EVIDENCE_PREVIEW_UNAVAILABLE",
            "来源存在，但当前无法打开对应预览位置",
        )
    result = success(location, "Evidence 原文定位成功")
    record_trace(
        session_id=trace_session,
        event_type="evidence_location",
        tool_name="locate_evidence",
        arguments={"evidence_id": clean_id},
        result=result,
        duration_ms=max(0, int((time.perf_counter() - started) * 1000)),
        result_status="success",
        metrics={
            "evidence_id": clean_id,
            "location_type": location.get("location_type"),
            "page_no": location.get("page_no"),
            "has_bbox": location.get("bbox") is not None,
        },
    )
    return result


def _locate_document(
    evidence: dict[str, Any], record: dict[str, Any]
) -> dict[str, Any]:
    chunks = database.get_document_chunks(record["file_id"])
    chunk = next(
        (
            item
            for item in chunks
            if item.get("chunk_id") == evidence.get("chunk_id")
            or (
                evidence.get("block_id")
                and (item.get("metadata") or {}).get("block_id")
                == evidence.get("block_id")
            )
        ),
        None,
    )
    region = _matching_region(record["file_id"], evidence.get("region_id"))
    if chunk is None and region is None:
        raise ValueError("Evidence 对应的活动索引或视觉 region 已不可用")
    metadata = (chunk or {}).get("metadata") or {}
    block = metadata.get("block") or {}
    page_no = (
        evidence.get("page_no") or (chunk or {}).get("page_no")
        or (region or {}).get("page_no")
    )
    bbox = (
        evidence.get("bbox") or metadata.get("bbox") or block.get("bbox")
        or (region or {}).get("bbox")
    )
    confidence = (
        evidence.get("confidence")
        if evidence.get("confidence") is not None
        else metadata.get("confidence", block.get("confidence"))
    )
    if confidence is None:
        confidence = (region or {}).get("confidence")
    base = {
        "evidence_id": evidence["evidence_id"],
        "source_type": evidence["source_type"],
        "file_id": record["file_id"],
        "file_name": record["file_name"],
        "block_id": evidence.get("block_id") or metadata.get("block_id"),
        "chunk_id": (chunk or {}).get("chunk_id"),
        "region_id": evidence.get("region_id") or (region or {}).get("region_id"),
        "text": (
            (chunk or {}).get("chunk_text") or block.get("content")
            or (region or {}).get("text") or evidence.get("text_excerpt") or ""
        ),
        "bbox": bbox,
        "confidence": confidence,
        "table_id": evidence.get("table") or metadata.get("table_id"),
        "cell_id": evidence.get("cell") or metadata.get("cell_id"),
        "row_index": (
            evidence.get("row_index")
            if evidence.get("row_index") is not None else metadata.get("row_index")
        ),
        "column_index": (
            evidence.get("column_index")
            if evidence.get("column_index") is not None else metadata.get("column_index")
        ),
        "line_number": evidence.get("line_number") or metadata.get("line_number"),
        "json_path": evidence.get("json_path") or metadata.get("json_path"),
        "slide_number": evidence.get("slide_number") or metadata.get("slide_number") or page_no,
        "recognition_type": evidence.get("recognition_type") or (region or {}).get("recognition_type"),
        "handwriting_confidence": evidence.get("handwriting_confidence"),
        "review_required": bool(evidence.get("review_required")),
        "conflict_sources": list(evidence.get("conflict_sources") or []),
        "evidence_status": evidence.get("evidence_status", "supported"),
    }
    if record["file_type"] == "pdf":
        if page_no is None:
            raise ValueError("PDF Evidence 缺少页码")
        return {
            **base,
            "location_type": "pdf",
            "page_no": int(page_no),
            "highlight": {"bbox": bbox} if bbox is not None else None,
        }
    if record["file_type"] == "presentation":
        return {**base, "location_type": "presentation", "slide_number": base["slide_number"]}
    if record["file_type"] == "txt":
        return {**base, "location_type": "txt", "line_number": base["line_number"]}
    if record["file_type"] == "json":
        return {**base, "location_type": "json", "json_path": base["json_path"]}
    return {
        **base,
        "location_type": "word",
        "paragraph_no": metadata.get("paragraph_no")
        or metadata.get("paragraph_start"),
        "word_region": metadata.get("word_region"),
    }


def _locate_image(
    evidence: dict[str, Any], record: dict[str, Any], path: Any
) -> dict[str, Any]:
    region = _matching_region(record["file_id"], evidence.get("region_id"))
    if evidence.get("region_id") and region is None:
        raise ValueError("Evidence 对应的视觉 region 已不可用")
    bbox = evidence.get("bbox") or (region or {}).get("bbox")
    if bbox is None:
        raise ValueError("图片 Evidence 缺少 bbox")
    preview = build_image_preview(path)
    if not preview["ok"]:
        raise ValueError("图片预览不可用")
    confidence = evidence.get("confidence")
    if confidence is None:
        confidence = (region or {}).get("confidence")
    recognition_type = evidence.get("recognition_type") or (region or {}).get(
        "recognition_type"
    )
    handwriting_confidence = evidence.get("handwriting_confidence")
    if handwriting_confidence is None and recognition_type == "handwritten":
        handwriting_confidence = confidence
    return {
        "evidence_id": evidence["evidence_id"],
        "location_type": "image",
        "source_type": evidence["source_type"],
        "file_id": record["file_id"],
        "file_name": record["file_name"],
        "image_no": int(evidence.get("image_no") or (region or {}).get("image_no") or 1),
        "page_no": evidence.get("page_no") or (region or {}).get("page_no"),
        "block_id": evidence.get("block_id"),
        "region_id": evidence.get("region_id") or (region or {}).get("region_id"),
        "text": (region or {}).get("text") or evidence.get("text_excerpt") or evidence.get("value_summary") or "",
        "bbox": bbox,
        "confidence": confidence,
        "recognition_type": recognition_type,
        "handwriting_confidence": handwriting_confidence,
        "review_required": bool(evidence.get("review_required")),
        "conflict_sources": list(evidence.get("conflict_sources") or []),
        "evidence_status": evidence.get("evidence_status", "supported"),
        "highlight": {"bbox": bbox},
        "preview": preview["data"],
    }


def _locate_visual_table(
    evidence: dict[str, Any], record: dict[str, Any], path: Any
) -> dict[str, Any]:
    from backend.services.visual_table import extract_visual_tables

    extracted = extract_visual_tables(record["file_id"])
    if not extracted["ok"]:
        raise ValueError("视觉表格当前不可用")
    table = next(
        (item["table"] for item in extracted["data"]["tables"]
         if item["table"]["table_id"] == evidence.get("table")),
        None,
    )
    cell = next(
        (item for item in (table or {}).get("cells", [])
         if item["cell_id"] == evidence.get("cell")),
        None,
    )
    if table is None or cell is None or cell.get("bbox") is None:
        raise ValueError("视觉表格 Cell 已不可用")
    base = {
        "evidence_id": evidence["evidence_id"],
        "source_type": evidence["source_type"],
        "file_id": record["file_id"],
        "file_name": record["file_name"],
        "page_no": cell.get("page_no"),
        "image_no": cell.get("image_no"),
        "table_id": table["table_id"],
        "cell_id": cell["cell_id"],
        "row_index": cell["row_index"],
        "column_index": cell["column_index"],
        "region_id": cell.get("source_region_id"),
        "text": cell.get("cell_text") or "",
        "bbox": cell["bbox"],
        "confidence": cell.get("confidence"),
        "recognition_type": cell.get("recognition_type"),
        "highlight": {"bbox": cell["bbox"]},
        "visual_table": True,
    }
    if record["file_type"] == "image":
        preview = build_image_preview(path)
        if not preview["ok"]:
            raise ValueError("图片预览不可用")
        return {**base, "location_type": "image", "preview": preview["data"]}
    if record["file_type"] == "pdf":
        return {**base, "location_type": "pdf"}
    raise ValueError("视觉表格来源类型不受支持")


def _matching_region(file_id: str, region_id: Any) -> dict[str, Any] | None:
    if not region_id:
        return None
    return next(
        (region for region in database.get_document_ocr_regions(file_id)
         if region.get("region_id") == region_id),
        None,
    )


def _locate_excel(
    evidence: dict[str, Any], record: dict[str, Any], path: Any
) -> dict[str, Any]:
    sheet_name = str(evidence.get("sheet") or evidence.get("table") or "")
    cell_id = str(evidence.get("cell") or "")
    if not sheet_name or not cell_id:
        raise ValueError("Excel Evidence 缺少 Sheet 或 Cell")
    row_index, column_index = coordinate_to_tuple(cell_id)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError("Evidence Sheet 已不存在")
        worksheet = workbook[sheet_name]
        if row_index > worksheet.max_row or column_index > worksheet.max_column:
            raise ValueError("Evidence Cell 已不存在")
        schema_rows = database.get_table_schema_records(record["file_id"], sheet_name)
        header_row = int(schema_rows[0]["header_row"]) if schema_rows else 1
        headers = [
            str(worksheet.cell(header_row, column).value or f"Column {column}")
            for column in range(1, worksheet.max_column + 1)
        ]
        row_values = [
            worksheet.cell(row_index, column).value
            for column in range(1, worksheet.max_column + 1)
        ]
        target_value = worksheet.cell(row_index, column_index).value
    finally:
        workbook.close()
    return {
        "evidence_id": evidence["evidence_id"],
        "location_type": "excel",
        "source_type": evidence["source_type"],
        "file_id": record["file_id"],
        "file_name": record["file_name"],
        "sheet": sheet_name,
        "table_id": evidence.get("table") or sheet_name,
        "cell_id": cell_id,
        "row_index": row_index,
        "column_index": column_index,
        "record_identifier": evidence.get("record_key"),
        "field": evidence.get("field"),
        "bbox": evidence.get("bbox"),
        "confidence": evidence.get("confidence"),
        "headers": headers,
        "row_values": row_values,
        "target_value": target_value,
    }


def _locate_csv(
    evidence: dict[str, Any], record: dict[str, Any], path: Any
) -> dict[str, Any]:
    table = parse_csv_table(path)
    row_number = int(evidence.get("row_number") or 0)
    column_name = str(evidence.get("column_name") or evidence.get("field") or "")
    if row_number < 2 or row_number - 2 >= len(table.rows):
        raise ValueError("CSV Evidence 行已不存在")
    if column_name not in table.headers:
        raise ValueError("CSV Evidence 列已不存在")
    column_number = table.headers.index(column_name) + 1
    row_values = table.rows[row_number - 2]
    return {
        "evidence_id": evidence["evidence_id"],
        "location_type": "csv",
        "source_type": evidence["source_type"],
        "file_id": record["file_id"],
        "file_name": record["file_name"],
        "table_id": evidence.get("table") or "CSV",
        "cell_id": evidence.get("cell"),
        "row_number": row_number,
        "column_number": column_number,
        "column_name": column_name,
        "record_identifier": evidence.get("record_key"),
        "headers": list(table.headers),
        "row_values": row_values,
        "target_value": row_values[column_number - 1],
    }


def _location_failure(
    evidence_id: str, session_id: str, error_code: str, message: str
) -> dict[str, Any]:
    result = failure(error_code, message)
    record_trace(
        session_id=session_id,
        event_type="evidence_location",
        tool_name="locate_evidence",
        arguments={"evidence_id": evidence_id},
        result=result,
        result_status="failed",
        error_code=error_code,
    )
    return result
