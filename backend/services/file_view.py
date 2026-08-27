"""Read-only file presentation metadata composed from existing repositories."""

from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook

from backend import database
from backend.services import ocr_service
from backend.services.file_lifecycle import get_file_lifecycle
from backend.services.file_locator import FileLocatorError, resolve_by_file_id
from backend.tools.excel_utils import failure, success
from backend.tools.file_tools import list_files


PREVIEW_MAX_SHEETS = 3
PREVIEW_MAX_ROWS = 6
PREVIEW_MAX_COLUMNS = 12
PREVIEW_MAX_TEXT_ITEMS = 10
PREVIEW_MAX_ITEM_CHARS = 500


def list_file_views(
    *,
    search: str | None = None,
    file_type: str | None = None,
    lifecycle_status: str | None = None,
    sort_by: str | None = None,
    sort_order: str = "desc",
) -> dict[str, Any]:
    """Return workspace files after deterministic Python filtering and sorting."""
    result = list_files()
    if not result["ok"]:
        return result
    items = [_enrich(item) for item in result["data"]]
    clean_search = (search or "").strip().casefold()
    if clean_search:
        items = [
            item for item in items
            if clean_search in str(item.get("file_name") or "").casefold()
        ]
    if file_type:
        items = [item for item in items if item.get("file_type") == file_type]
    if lifecycle_status:
        items = [
            item for item in items
            if item.get("lifecycle_status") == lifecycle_status
        ]
    if sort_by:
        items.sort(
            key=lambda item: _sort_value(item, sort_by),
            reverse=sort_order == "desc",
        )
    return success(items, f"找到 {len(items)} 个受支持文件")


def get_file_view(file_id: str) -> dict[str, Any]:
    clean = str(file_id).strip()
    record = database.get_file_record_by_id(clean)
    if record is None:
        return failure("FILE_NOT_FOUND", "未找到指定文件")
    listed = list_files()
    item = next(
        (entry for entry in (listed.get("data") or []) if entry.get("file_id") == clean),
        None,
    )
    if item is None:
        return failure("FILE_NOT_FOUND", "文件记录存在，但物理文件不可用")
    return success(_enrich(item), "文件展示信息读取成功")


def get_file_preview(file_id: str) -> dict[str, Any]:
    """Return a bounded, read-only preview resolved only by stable file ID."""
    detail = get_file_view(file_id)
    if not detail["ok"]:
        return detail
    item = detail["data"]
    try:
        path = resolve_by_file_id(item["file_id"])
        if item["file_type"] == "excel":
            preview = _preview_excel(path)
        elif item["file_type"] == "word":
            preview = _preview_word(item["file_id"], path)
        elif item["file_type"] == "pdf":
            preview = _preview_pdf(item["file_id"], path)
        else:
            return failure("UNSUPPORTED_FILE_TYPE", "当前文件类型不支持快速预览")
    except FileLocatorError as exc:
        return failure("FILE_NOT_FOUND", str(exc))
    except Exception:
        return failure("FILE_PREVIEW_ERROR", "文件预览生成失败")
    return success(
        {
            "file_id": item["file_id"],
            "file_name": item["file_name"],
            "file_type": item["file_type"],
            **preview,
        },
        "文件预览读取成功",
    )


def _enrich(item: dict[str, Any]) -> dict[str, Any]:
    enriched = dict(item)
    record = (
        database.get_file_record_by_id(item["file_id"])
        if item.get("file_id")
        else database.get_file_record(item["file_name"])
    )
    enriched["uploaded_at"] = record.get("created_at") if record else None
    enriched["created_time"] = record.get("created_at") if record else None
    enriched["updated_at"] = record.get("updated_at") if record else None
    enriched["status"] = record.get("status") if record else "active"
    enriched["writable"] = bool(record.get("writable")) if record else False
    enriched["deletable"] = bool(record.get("deletable")) if record else False
    lifecycle = get_file_lifecycle(record["file_id"]) if record else None
    lifecycle_data = lifecycle.get("data") if lifecycle and lifecycle.get("ok") else {}
    enriched["canonical_status"] = lifecycle_data.get("status")
    enriched["error_summary"] = lifecycle_data.get("error")
    enriched["schema_summary"] = None
    if item.get("file_type") == "excel" and item.get("file_id"):
        schemas = database.get_table_schema_records(item["file_id"])
        enriched["schema_summary"] = {
            "sheet_count": len(schemas),
            "field_count": sum(int(schema.get("column_count") or 0) for schema in schemas),
            "row_count": sum(int(schema.get("row_count") or 0) for schema in schemas),
            "sheets": [
                {
                    "sheet_name": schema["sheet_name"],
                    "field_count": schema["column_count"],
                    "row_count": schema["row_count"],
                }
                for schema in schemas
            ],
        } if schemas else None
    return enriched


def _preview_excel(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for worksheet in workbook.worksheets[:PREVIEW_MAX_SHEETS]:
            rows = []
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=min(worksheet.max_row, PREVIEW_MAX_ROWS),
                max_col=min(worksheet.max_column, PREVIEW_MAX_COLUMNS),
                values_only=True,
            ):
                rows.append([_preview_value(value) for value in row])
            sheets.append({"sheet_name": worksheet.title, "rows": rows})
        return {
            "preview_type": "table",
            "sheets": sheets,
            "truncated": len(workbook.sheetnames) > PREVIEW_MAX_SHEETS
            or any(
                sheet.max_row > PREVIEW_MAX_ROWS
                or sheet.max_column > PREVIEW_MAX_COLUMNS
                for sheet in workbook.worksheets[:PREVIEW_MAX_SHEETS]
            ),
        }
    finally:
        workbook.close()


def _preview_word(file_id: str, path: Path) -> dict[str, Any]:
    indexed = _indexed_text_preview(file_id)
    if indexed:
        return indexed
    document = Document(path)
    paragraphs = [
        paragraph.text.strip()[:PREVIEW_MAX_ITEM_CHARS]
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    return {
        "preview_type": "text",
        "items": [
            {"location": f"段落 {index}", "text": text}
            for index, text in enumerate(paragraphs[:PREVIEW_MAX_TEXT_ITEMS], start=1)
        ],
        "truncated": len(paragraphs) > PREVIEW_MAX_TEXT_ITEMS,
    }


def _preview_pdf(file_id: str, path: Path) -> dict[str, Any]:
    indexed = _indexed_text_preview(file_id)
    if indexed:
        return indexed
    detected = ocr_service.detect_pdf_text(path)
    if not detected["ok"]:
        raise ValueError("PDF preview unavailable")
    pages = detected["data"]["pages"]
    return {
        "preview_type": "text",
        "items": [
            {
                "location": f"第 {page['page']} 页",
                "text": str(page.get("text") or "")[:PREVIEW_MAX_ITEM_CHARS],
            }
            for page in pages[:PREVIEW_MAX_TEXT_ITEMS]
        ],
        "truncated": len(pages) > PREVIEW_MAX_TEXT_ITEMS,
    }


def _indexed_text_preview(file_id: str) -> dict[str, Any] | None:
    chunks = database.get_document_chunks(file_id)
    if not chunks:
        return None
    items = []
    for chunk in chunks[:PREVIEW_MAX_TEXT_ITEMS]:
        page_no = chunk.get("page_no")
        block = (chunk.get("metadata") or {}).get("block") or {}
        location = f"第 {page_no} 页" if page_no else f"Block {block.get('block_id') or chunk['chunk_id']}"
        items.append(
            {
                "location": location,
                "text": str(chunk.get("chunk_text") or "")[:PREVIEW_MAX_ITEM_CHARS],
                "block_id": block.get("block_id"),
                "page_no": page_no,
            }
        )
    return {
        "preview_type": "text",
        "items": items,
        "truncated": len(chunks) > PREVIEW_MAX_TEXT_ITEMS,
    }


def _preview_value(value: Any) -> str:
    if value is None:
        return ""
    return str(value)[:PREVIEW_MAX_ITEM_CHARS]


def _sort_value(item: dict[str, Any], sort_by: str) -> Any:
    if sort_by == "size":
        return int(item.get("size") or 0)
    return str(item.get("created_time") or "")
