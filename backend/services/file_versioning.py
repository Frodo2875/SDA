"""Word diff previews and immutable snapshots around the existing HITL action."""

import hashlib
import json
import os
import tempfile
from pathlib import Path
from typing import Any, Callable, Literal
from uuid import uuid4

from docx import Document
from pydantic import BaseModel, ConfigDict, Field, model_validator

from backend import database
from backend.services.file_locator import FileLocatorError, resolve_by_file_id
from backend.tools.excel_utils import failure, success


class WordDiffOperation(BaseModel):
    """Validated operation frozen into a pending action."""

    model_config = ConfigDict(extra="forbid")

    operation_type: Literal["append", "undo", "rollback"]
    content: str | None = Field(default=None, max_length=100_000)
    version_id: str | None = Field(default=None, max_length=128)

    @model_validator(mode="after")
    def validate_operation(self) -> "WordDiffOperation":
        if self.operation_type == "append" and not (self.content or "").strip():
            raise ValueError("append 操作必须提供非空 content")
        if self.operation_type == "rollback" and not (self.version_id or "").strip():
            raise ValueError("rollback 操作必须提供 version_id")
        if self.operation_type == "undo" and self.version_id is not None:
            raise ValueError("undo 操作不接受 version_id")
        return self


WriteHandler = Callable[[str, str], dict[str, Any]]


def preview_word_diff(
    file_id: str,
    operation: WordDiffOperation | dict[str, Any],
) -> dict[str, Any]:
    """Build a deterministic preview without modifying the document or versions."""
    record, path, error = _writable_word(file_id)
    if error is not None:
        return error
    try:
        validated = (
            operation
            if isinstance(operation, WordDiffOperation)
            else WordDiffOperation.model_validate(operation)
        )
    except Exception as exc:
        return failure("INVALID_DIFF_OPERATION", f"Diff 操作参数无效：{exc}")

    try:
        before_bytes = path.read_bytes()
        before_text = _word_text(path)
    except Exception:
        return failure("WORD_READ_ERROR", "读取 Word 生成 Diff 失败")
    target_version = None
    if validated.operation_type == "append":
        after_text = "\n".join(
            item for item in (before_text, validated.content.strip()) if item
        )
        location = "文档末尾"
        target_object = "Word 文档正文"
        impact_scope = "新增 1 个段落"
    else:
        target_version, version_error = _target_version(
            record["file_id"], validated.operation_type, validated.version_id
        )
        if version_error is not None:
            return version_error
        snapshot_path, snapshot_error = _snapshot_path(target_version)
        if snapshot_error is not None:
            return snapshot_error
        try:
            after_text = _word_text(snapshot_path)
        except Exception:
            return failure("VERSION_READ_ERROR", "读取目标版本生成 Diff 失败")
        location = "整个 Word 文档"
        target_object = f"版本 {target_version['version_number']}"
        impact_scope = "恢复整个文档内容"

    normalized_operation = validated.model_dump()
    normalized_operation["expected_hash"] = _hash_bytes(before_bytes)
    normalized_operation["target_version_id"] = (
        target_version["version_id"] if target_version else None
    )
    diff = {
        "target_file": record["file_name"],
        "target_object": target_object,
        "location": location,
        "operation_type": validated.operation_type,
        "before": before_text,
        "after": after_text,
        "impact_scope": impact_scope,
    }
    return success(
        {
            **diff,
            "file_id": record["file_id"],
            "operation": normalized_operation,
        },
        "Diff 预览已生成，实际文件未修改",
    )


def list_versions(file_id: str) -> dict[str, Any]:
    record = database.get_file_record_by_id(str(file_id).strip())
    if record is None:
        return failure("FILE_NOT_FOUND", "未找到指定文件")
    return success(database.list_file_versions(record["file_id"]), "文件版本读取成功")


def execute_versioned_word_action(
    action: dict[str, Any],
    *,
    write_handler: WriteHandler,
) -> dict[str, Any]:
    """Execute one frozen Word action and compensate every failure before returning."""
    operation = _operation_from_action(action)
    file_id = str(action.get("file_id") or "").strip()
    if not file_id:
        record = database.get_file_record(action["target_file"])
        file_id = str(record["file_id"]) if record else ""
    record, path, error = _writable_word(file_id)
    if error is not None:
        return error

    original_bytes = path.read_bytes()
    expected_hash = operation.get("expected_hash")
    if expected_hash and expected_hash != _hash_bytes(original_bytes):
        return failure(
            "FILE_CHANGED_SINCE_PREVIEW",
            "文件在 Diff 预览后已发生变化，请重新生成预览并确认",
        )

    before_id = uuid4().hex
    after_id = uuid4().hex
    before_path = _version_storage_path(file_id, before_id)
    after_path = _version_storage_path(file_id, after_id)
    try:
        _write_snapshot(before_path, original_bytes)
        action_result = _apply_operation(
            path=path,
            record=record,
            operation=operation,
            write_handler=write_handler,
        )
        if not action_result["ok"]:
            _restore_if_changed(path, original_bytes)
            _cleanup_paths(before_path, after_path)
            return action_result

        Document(path)
        updated_bytes = path.read_bytes()
        _write_snapshot(after_path, updated_bytes)
        timestamp = database.utc_now()
        transition = database.commit_file_version_transition(
            file_id=file_id,
            before_version=_version_metadata(
                version_id=before_id,
                file_id=file_id,
                path=before_path,
                content=original_bytes,
                action=action,
                change_type="snapshot",
                created_at=timestamp,
            ),
            after_version=_version_metadata(
                version_id=after_id,
                file_id=file_id,
                path=after_path,
                content=updated_bytes,
                action=action,
                change_type=str(operation["operation_type"]),
                created_at=timestamp,
            ),
        )
        if transition["reused_before"]:
            _cleanup_paths(before_path)
        return success(
            {
                **(action_result.get("data") or {}),
                "file_id": file_id,
                "before_version": transition["before_version"],
                "new_version": transition["new_version"],
            },
            "Word 操作已安全执行并创建新版本",
        )
    except Exception:
        restore_error = _restore_if_changed(path, original_bytes)
        _cleanup_paths(before_path, after_path)
        if restore_error is not None:
            return failure(
                "WORD_RECOVERY_ERROR",
                "版本写入失败且原文件恢复失败，需要人工检查文件状态",
            )
        return failure(
            "VERSION_WRITE_ERROR",
            "版本化 Word 操作失败，原文件保持不变且未留下版本记录",
        )


def _operation_from_action(action: dict[str, Any]) -> dict[str, Any]:
    raw = action.get("operation_json")
    if raw:
        parsed = json.loads(raw)
        WordDiffOperation.model_validate(
            {
                "operation_type": parsed.get("operation_type"),
                "content": parsed.get("content"),
                "version_id": parsed.get("version_id"),
            }
        )
        return parsed
    return {
        "operation_type": "append",
        "content": action["content"],
        "version_id": None,
        "target_version_id": None,
        "expected_hash": None,
    }


def _apply_operation(
    *,
    path: Path,
    record: dict[str, Any],
    operation: dict[str, Any],
    write_handler: WriteHandler,
) -> dict[str, Any]:
    if operation["operation_type"] == "append":
        return write_handler(record["file_name"], str(operation["content"]))
    target_id = operation.get("target_version_id") or operation.get("version_id")
    target = database.get_file_version(str(target_id)) if target_id else None
    if target is None or target["file_id"] != record["file_id"]:
        return failure("VERSION_NOT_FOUND", "待恢复版本不存在或不属于目标文件")
    snapshot, error = _snapshot_path(target)
    if error is not None:
        return error
    try:
        _atomic_replace_bytes(path, snapshot.read_bytes())
    except Exception:
        return failure("WORD_ROLLBACK_ERROR", "Word 恢复失败，当前文件保持不变")
    return success(
        {"file_name": record["file_name"], "restored_version_id": target["version_id"]},
        "Word 内容已恢复到指定版本",
    )


def _target_version(
    file_id: str, operation_type: str, version_id: str | None
) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    if operation_type == "undo":
        latest = database.get_latest_file_version(file_id)
        if latest is None or not latest.get("parent_version_id"):
            return None, failure("NO_UNDO_VERSION", "当前文件没有可撤销的最近修改")
        target = database.get_file_version(latest["parent_version_id"])
    else:
        target = database.get_file_version(str(version_id or "").strip())
    if target is None or target["file_id"] != file_id or target["status"] != "available":
        return None, failure("VERSION_NOT_FOUND", "指定版本不存在或不属于目标文件")
    return target, None


def _writable_word(
    file_id: str,
) -> tuple[dict[str, Any] | None, Path | None, dict[str, Any] | None]:
    clean_id = str(file_id or "").strip()
    record = database.get_file_record_by_id(clean_id)
    if record is None:
        return None, None, failure("FILE_NOT_FOUND", "未找到指定文件")
    if record["file_type"] != "word":
        return None, None, failure("INVALID_TARGET_FILE", "目标文件必须是 Word 文档")
    if not bool(record["writable"]):
        return None, None, failure("FILE_WRITE_FORBIDDEN", "该系统文件为只读，禁止修改")
    if record["lifecycle_status"] != "ready":
        return None, None, failure("FILE_NOT_READY", "目标文件当前不可写入")
    try:
        path = resolve_by_file_id(clean_id)
    except FileLocatorError as exc:
        return None, None, failure("FILE_NOT_FOUND", str(exc))
    if path.suffix.lower() != ".docx":
        return None, None, failure("INVALID_TARGET_FILE", "目标文件必须是 Word 文档")
    return record, path, None


def _word_text(path: Path) -> str:
    document = Document(path)
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)


def _snapshot_path(
    version: dict[str, Any],
) -> tuple[Path | None, dict[str, Any] | None]:
    raw = str(version.get("storage_path") or "")
    prefix = "data/versions/"
    if not raw.startswith(prefix) or ".." in Path(raw).parts:
        return None, failure("INVALID_VERSION_PATH", "版本存储路径无效")
    relative = Path(raw[len("data/versions/"):])
    versions_root = (database.DB_PATH.parent / "versions").resolve()
    path = (versions_root / relative).resolve()
    if versions_root not in path.parents or not path.is_file():
        return None, failure("VERSION_FILE_MISSING", "版本快照文件不存在")
    try:
        content = path.read_bytes()
    except OSError:
        return None, failure("VERSION_READ_ERROR", "版本快照读取失败")
    if _hash_bytes(content) != version["content_hash"]:
        return None, failure("VERSION_HASH_MISMATCH", "版本快照完整性校验失败")
    return path, None


def _version_storage_path(file_id: str, version_id: str) -> Path:
    if not all(char.isalnum() or char in {"-", "_"} for char in file_id):
        raise ValueError("非法 file_id")
    return database.DB_PATH.parent / "versions" / file_id / f"{version_id}.docx"


def _write_snapshot(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb", dir=path.parent, prefix=".version-", suffix=".tmp", delete=False
        ) as output:
            output.write(content)
            output.flush()
            os.fsync(output.fileno())
            temporary = Path(output.name)
        os.replace(temporary, path)
        temporary = None
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def _atomic_replace_bytes(path: Path, content: bytes) -> None:
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb", dir=path.parent, prefix=f".{path.stem}-restore-",
            suffix=".docx", delete=False
        ) as output:
            output.write(content)
            output.flush()
            os.fsync(output.fileno())
            temporary = Path(output.name)
        Document(temporary)
        os.replace(temporary, path)
        temporary = None
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def _restore_if_changed(path: Path, original: bytes) -> Exception | None:
    try:
        if path.read_bytes() != original:
            _atomic_replace_bytes(path, original)
        return None
    except Exception as exc:
        return exc


def _version_metadata(
    *,
    version_id: str,
    file_id: str,
    path: Path,
    content: bytes,
    action: dict[str, Any],
    change_type: str,
    created_at: str,
) -> dict[str, Any]:
    relative = path.relative_to(database.DB_PATH.parent / "versions").as_posix()
    return {
        "version_id": version_id,
        "storage_path": f"data/versions/{relative}",
        "content_hash": _hash_bytes(content),
        "size": len(content),
        "task_id": action.get("task_id"),
        "session_id": action["session_id"],
        "change_type": change_type,
        "created_at": created_at,
        "status": "available",
    }


def _hash_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _cleanup_paths(*paths: Path) -> None:
    for path in paths:
        try:
            path.unlink(missing_ok=True)
        except OSError:
            pass
