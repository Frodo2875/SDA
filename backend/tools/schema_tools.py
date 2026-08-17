"""Deterministic Excel Schema Discovery tools without LLM involvement."""

import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from backend import database
from backend.services.file_locator import FileLocatorError, resolve_by_file_id
from backend.services.field_semantics import map_field_semantics
from backend.tools.excel_utils import failure, success


FILE_ID_PATTERN = re.compile(r"^[0-9a-fA-F]{32}$")
MAX_HEADER_SCAN_ROWS = 50
MAX_DISCOVERY_COLUMNS = 200
MIN_HEADER_CONFIDENCE = 0.78


def inspect_excel(file_id: str) -> dict[str, Any]:
    """Discover and persist independent schemas for every Sheet in an Excel file."""
    clean_file_id, validation_error = _validate_file_id(file_id)
    if validation_error is not None:
        return validation_error
    record = database.get_file_record_by_id(clean_file_id)
    if record is None:
        return failure("FILE_NOT_FOUND", "未找到指定 file_id 的文件")
    if record["file_type"] != "excel":
        return failure("UNSUPPORTED_FILE_TYPE", "inspect_excel 仅支持 .xlsx 文件")
    if record["lifecycle_status"] == "deleted":
        return failure("FILE_NOT_FOUND", "文件已经删除")

    state_started = database.update_file_state(
        file_id=clean_file_id,
        lifecycle_status="processing",
        parse_status="processing",
        queryable=False,
        index_status="not_required",
    )
    if not state_started:
        return failure("FILE_STATE_ERROR", "无法将文件切换到 processing 状态")

    try:
        path = resolve_by_file_id(clean_file_id)
    except FileLocatorError as exc:
        return _fail_inspection(clean_file_id, "FILE_NOT_FOUND", str(exc))

    workbook = None
    try:
        workbook = load_workbook(
            path,
            read_only=False,
            data_only=True,
            keep_vba=False,
            keep_links=False,
        )
        schemas = [_inspect_worksheet(worksheet) for worksheet in workbook.worksheets]
    except Exception:
        return _fail_inspection(
            clean_file_id,
            "EXCEL_PARSE_ERROR",
            "Excel 文件无法正常解析，文件可能已损坏",
        )
    finally:
        if workbook is not None:
            workbook.close()

    try:
        database.replace_table_schemas(clean_file_id, schemas)
    except Exception:
        return _fail_inspection(
            clean_file_id,
            "SCHEMA_PERSIST_ERROR",
            "Excel Schema 保存失败",
            clear_schemas=False,
        )

    failed_sheets = [
        schema for schema in schemas if schema["detection_status"] != "detected"
    ]
    if failed_sheets:
        state_failed = database.update_file_state(
            file_id=clean_file_id,
            lifecycle_status="failed",
            parse_status="failed",
            queryable=False,
            index_status="not_required",
        )
        if not state_failed:
            return failure("FILE_STATE_ERROR", "Schema 已保存，但文件失败状态更新失败")
        reasons = [
            f"{schema['sheet_name']}：{schema['detection_message']}"
            for schema in failed_sheets
        ]
        return failure(
            "HEADER_DETECTION_FAILED",
            "存在无法可靠识别表头的 Sheet：" + "；".join(reasons),
            _inspection_payload(record, schemas),
        )

    state_ready = database.update_file_state(
        file_id=clean_file_id,
        lifecycle_status="ready",
        parse_status="parsed",
        queryable=True,
        index_status="not_required",
    )
    if not state_ready:
        return failure("FILE_STATE_ERROR", "Schema 已保存，但文件就绪状态更新失败")
    stored = database.get_table_schema_records(clean_file_id)
    return success(
        _inspection_payload(record, [_format_schema(schema) for schema in stored]),
        f"已识别 {len(stored)} 个 Sheet 的表结构",
    )


def get_table_schema(
    file_id: str,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    """Return persisted Schema Discovery output without opening Excel again."""
    clean_file_id, validation_error = _validate_file_id(file_id)
    if validation_error is not None:
        return validation_error
    clean_sheet_name = None
    if sheet_name is not None:
        if not isinstance(sheet_name, str) or not sheet_name.strip():
            return failure("INVALID_SHEET_NAME", "sheet_name 不能为空")
        clean_sheet_name = sheet_name.strip()

    record = database.get_file_record_by_id(clean_file_id)
    if record is None:
        return failure("FILE_NOT_FOUND", "未找到指定 file_id 的文件")
    try:
        schemas = database.get_table_schema_records(clean_file_id, clean_sheet_name)
    except Exception:
        return failure("SCHEMA_READ_ERROR", "读取已保存 Schema 失败")
    if not schemas:
        error_code = "SHEET_SCHEMA_NOT_FOUND" if clean_sheet_name else "SCHEMA_NOT_FOUND"
        message = (
            f"未找到 Sheet 的已保存 Schema：{clean_sheet_name}"
            if clean_sheet_name
            else "该文件尚无已保存 Schema，请先执行 inspect_excel"
        )
        return failure(error_code, message)
    formatted = [_format_schema(schema) for schema in schemas]
    return success(
        _inspection_payload(record, formatted),
        "已读取真实保存的 Excel Schema",
    )


def _inspect_worksheet(worksheet: Worksheet) -> dict[str, Any]:
    header = _detect_header(worksheet)
    if not header["ok"]:
        return {
            "sheet_name": worksheet.title,
            "header_row": None,
            "data_start_row": None,
            "row_count": 0,
            "column_count": 0,
            "detection_status": "failed",
            "confidence": 0.0,
            "detection_message": header["message"],
            "fields": [],
        }

    header_row = header["row"]
    columns = header["columns"]
    names = header["names"]
    data_rows = []
    for row_index in range(header_row + 1, worksheet.max_row + 1):
        values = [worksheet.cell(row=row_index, column=index).value for index in columns]
        if any(not _is_empty(value) for value in values):
            data_rows.append(values)

    fields = []
    row_count = len(data_rows)
    for position, (column_index, source_name) in enumerate(zip(columns, names)):
        values = [row[position] for row in data_rows]
        non_null = [value for value in values if not _is_empty(value)]
        null_count = row_count - len(non_null)
        semantic = map_field_semantics(source_name)
        fields.append(
            {
                "source_name": source_name,
                "source_index": column_index,
                "inferred_type": _infer_type(non_null),
                "nullable": null_count > 0,
                "null_count": null_count,
                "null_ratio": round(null_count / row_count, 6) if row_count else 0.0,
                "unique_count": len({_unique_key(value) for value in non_null}),
                "semantic_type": semantic["semantic_type"],
                "confidence": semantic["mapping_confidence"],
                "sensitive": semantic["sensitive"],
                "canonical_name": semantic["canonical_name"],
                "mapping_confidence": semantic["mapping_confidence"],
                "mapping_source": semantic["mapping_source"],
            }
        )

    return {
        "sheet_name": worksheet.title,
        "header_row": header_row,
        "data_start_row": header_row + 1,
        "row_count": row_count,
        "column_count": len(columns),
        "detection_status": "detected",
        "confidence": header["confidence"],
        "detection_message": "表头识别成功",
        "fields": fields,
    }


def _detect_header(worksheet: Worksheet) -> dict[str, Any]:
    if worksheet.max_row == 1 and worksheet.max_column == 1:
        if _is_empty(worksheet.cell(row=1, column=1).value):
            return {"ok": False, "message": "Sheet 为空，无法识别表头"}
    if worksheet.max_column > MAX_DISCOVERY_COLUMNS:
        return {
            "ok": False,
            "message": f"字段数超过安全识别上限 {MAX_DISCOVERY_COLUMNS}",
        }

    max_row = min(worksheet.max_row, MAX_HEADER_SCAN_ROWS)
    max_column = min(worksheet.max_column, MAX_DISCOVERY_COLUMNS)
    candidates = []
    invalid_header_rows = []
    for row_index in range(1, max_row + 1):
        row_values = [
            worksheet.cell(row=row_index, column=column).value
            for column in range(1, max_column + 1)
        ]
        nonempty_columns = [
            index + 1 for index, value in enumerate(row_values) if not _is_empty(value)
        ]
        if len(nonempty_columns) < 2:
            continue
        first_column, last_column = nonempty_columns[0], nonempty_columns[-1]
        span = list(range(first_column, last_column + 1))
        values = [worksheet.cell(row=row_index, column=index).value for index in span]
        string_ratio = sum(isinstance(value, str) for value in values) / len(values)
        duplicate_names = len({_header_name(value) for value in values}) != len(values)
        has_gaps = any(_is_empty(value) for value in values)
        merged = _row_has_merged_span(worksheet, row_index, first_column, last_column)
        next_values = (
            [worksheet.cell(row=row_index + 1, column=index).value for index in span]
            if row_index < worksheet.max_row
            else []
        )
        next_nonempty = sum(not _is_empty(value) for value in next_values)
        next_coverage = next_nonempty / len(span) if span else 0.0
        header_like = string_ratio >= 0.8 and next_coverage >= 0.5
        if header_like and (duplicate_names or has_gaps):
            invalid_header_rows.append((row_index, first_column, last_column))
        if duplicate_names or has_gaps or merged or string_ratio < 0.8:
            continue
        if next_coverage < 0.5:
            continue

        next_row_nonempty_columns = [
            column
            for column in range(1, max_column + 1)
            if not _is_empty(worksheet.cell(row=row_index + 1, column=column).value)
        ]
        if next_row_nonempty_columns and (
            next_row_nonempty_columns[0] < first_column
            or next_row_nonempty_columns[-1] > last_column
            or len(next_row_nonempty_columns) > len(span)
        ):
            continue

        contrast = (
            sum(
                not isinstance(value, str)
                for value in next_values
                if not _is_empty(value)
            )
            / next_nonempty
            if next_nonempty
            else 0.0
        )
        confidence = round(
            0.35 * string_ratio
            + 0.20
            + 0.15
            + 0.20 * next_coverage
            + 0.10 * contrast,
            4,
        )
        if confidence >= MIN_HEADER_CONFIDENCE:
            candidates.append(
                {
                    "row": row_index,
                    "columns": span,
                    "names": [_header_name(value) for value in values],
                    "confidence": confidence,
                }
            )

    if not candidates:
        return {"ok": False, "message": "未找到满足可靠性阈值的连续唯一表头"}
    candidates.sort(
        key=lambda item: (-item["confidence"], -len(item["columns"]), item["row"])
    )
    selected = candidates[0]
    for row_index, first_column, last_column in invalid_header_rows:
        if row_index < selected["row"] and selected["row"] - row_index <= 1:
            selected_first = selected["columns"][0]
            selected_last = selected["columns"][-1]
            if first_column == selected_first and last_column == selected_last:
                return {
                    "ok": False,
                    "message": "候选表头包含空字段或重复字段，拒绝强行推断",
                }
    return {"ok": True, **selected}


def _row_has_merged_span(
    worksheet: Worksheet,
    row_index: int,
    first_column: int,
    last_column: int,
) -> bool:
    return any(
        merged.min_row <= row_index <= merged.max_row
        and merged.max_col > merged.min_col
        and merged.max_col >= first_column
        and merged.min_col <= last_column
        for merged in worksheet.merged_cells.ranges
    )


def _infer_type(values: list[Any]) -> str:
    if not values:
        return "empty"
    kinds = {_value_type(value) for value in values}
    if kinds <= {"integer", "number"}:
        return "number" if "number" in kinds else "integer"
    if len(kinds) == 1:
        return next(iter(kinds))
    return "mixed"


def _value_type(value: Any) -> str:
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "number"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, time):
        return "time"
    if isinstance(value, str):
        return "string"
    return "other"


def _format_schema(schema: dict[str, Any]) -> dict[str, Any]:
    fields = []
    for raw_field in schema.get("fields", []):
        field = dict(raw_field)
        field["nullable"] = bool(field["nullable"])
        field["sensitive"] = bool(field["sensitive"])
        fields.append(field)
    formatted = dict(schema)
    formatted["fields"] = fields
    formatted["possible_unique_fields"] = [
        field["source_name"]
        for field in fields
        if formatted["row_count"] >= 2
        and field["null_count"] == 0
        and field["unique_count"] == formatted["row_count"]
    ]
    formatted["possible_name_fields"] = [
        field["source_name"] for field in fields if field["semantic_type"] == "name"
    ]
    formatted["possible_sensitive_fields"] = [
        field["source_name"] for field in fields if field["sensitive"]
    ]
    return formatted


def _inspection_payload(
    record: dict[str, Any],
    schemas: list[dict[str, Any]],
) -> dict[str, Any]:
    return {
        "file_id": record["file_id"],
        "file_name": record["file_name"],
        "sheet_count": len(schemas),
        "sheet_names": [schema["sheet_name"] for schema in schemas],
        "schemas": schemas,
    }


def _fail_inspection(
    file_id: str,
    error_code: str,
    message: str,
    *,
    clear_schemas: bool = True,
) -> dict[str, Any]:
    try:
        if clear_schemas:
            database.delete_table_schemas(file_id)
        database.update_file_state(
            file_id=file_id,
            lifecycle_status="failed",
            parse_status="failed",
            queryable=False,
            index_status="not_required",
        )
    except Exception:
        return failure(
            "SCHEMA_STATE_ERROR",
            f"{message}；同时无法安全更新解析状态",
            {"file_id": file_id},
        )
    return failure(error_code, message, {"file_id": file_id})


def _validate_file_id(
    file_id: str,
) -> tuple[str, dict[str, Any] | None]:
    if not isinstance(file_id, str) or not FILE_ID_PATTERN.fullmatch(file_id.strip()):
        return "", failure("INVALID_FILE_ID", "file_id 必须是 32 位十六进制标识")
    return file_id.strip(), None


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _header_name(value: Any) -> str:
    return str(value).strip()


def _unique_key(value: Any) -> tuple[str, str]:
    return type(value).__name__, str(value).strip() if isinstance(value, str) else str(value)
