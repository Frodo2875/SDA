"""Safe deterministic querying and aggregation over discovered Excel tables."""

from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from backend import database
from backend.evidence import build_evidence, make_evidence_id
from backend.services.file_locator import FileLocatorError, resolve_by_file_id
from backend.services.multiformat_parser import parse_csv_table
from backend.tool_models import (
    AggregateTableArguments,
    QueryTableArguments,
    TableFilter,
    TableOrder,
)
from backend.tools.excel_utils import failure, success


def query_table(
    file_id: str,
    sheet: str,
    filters: list[dict[str, Any]],
    select: list[str],
    order_by: dict[str, Any] | list[dict[str, Any]] | None = None,
    limit: int | None = None,
) -> dict[str, Any]:
    """Filter, select, sort, and limit one discovered Sheet without SQL or eval."""
    try:
        arguments = QueryTableArguments.model_validate(
            {
                "file_id": file_id,
                "sheet": sheet,
                "filters": filters,
                "select": select,
                "order_by": order_by,
                "limit": limit,
            }
        )
    except ValidationError as exc:
        return failure("INVALID_QUERY_ARGUMENTS", _validation_message(exc))

    context_result = load_table_data(arguments.file_id, arguments.sheet)
    if not context_result["ok"]:
        return context_result
    context = context_result["data"]
    orders = _normalize_orders(arguments.order_by)
    requested_fields = (
        arguments.select
        + [item.field for item in arguments.filters]
        + [item.field for item in orders]
    )
    field_error = _validate_fields(context["fields"], requested_fields)
    if field_error is not None:
        return field_error

    filtered_result = _apply_filters(
        context["rows"], arguments.filters, context["fields"]
    )
    if not filtered_result["ok"]:
        return filtered_result
    matched_rows = filtered_result["data"]
    ordered_rows = _apply_order(matched_rows, orders)
    returned_rows = ordered_rows[: arguments.limit] if arguments.limit else ordered_rows
    result_rows = [
        {field: _json_value(row.get(field)) for field in arguments.select}
        for row in returned_rows
    ]
    status = "found" if matched_rows else "not_found"
    warnings = _query_warnings(
        context["fields"], arguments.select, len(matched_rows), len(result_rows)
    )
    response = success(
        {
            "status": status,
            "rows": result_rows,
            "matched_rows": len(matched_rows),
            "returned_rows": len(result_rows),
        },
        "表格查询完成" if matched_rows else "查询条件未匹配到记录",
    )
    response.update(
        {
            "evidence": _evidence(
                context,
                fields=arguments.select,
                filters=arguments.filters,
                matched_rows=len(matched_rows),
            ),
            "warnings": warnings,
            "result_summary": (
                f"匹配 {len(matched_rows)} 行，返回 {len(result_rows)} 行"
            ),
            "evidence_chain": _structured_evidence(
                context, returned_rows, arguments.select
            ),
        }
    )
    return response


def aggregate_table(
    file_id: str,
    sheet: str,
    operation: str,
    field: str | None = None,
    group_by: str | list[str] | None = None,
    filters: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    """Calculate an allow-listed aggregate from real Excel cells in Python."""
    try:
        arguments = AggregateTableArguments.model_validate(
            {
                "file_id": file_id,
                "sheet": sheet,
                "operation": operation,
                "field": field,
                "group_by": group_by or [],
                "filters": filters or [],
            }
        )
    except ValidationError as exc:
        return failure("INVALID_AGGREGATE_ARGUMENTS", _validation_message(exc))

    context_result = load_table_data(arguments.file_id, arguments.sheet)
    if not context_result["ok"]:
        return context_result
    context = context_result["data"]
    requested_fields = list(arguments.group_by)
    if arguments.field is not None:
        requested_fields.append(arguments.field)
    requested_fields.extend(item.field for item in arguments.filters)
    field_error = _validate_fields(context["fields"], requested_fields)
    if field_error is not None:
        return field_error
    if arguments.operation in {"sum", "avg"}:
        inferred_type = context["fields"][arguments.field]["inferred_type"]
        if inferred_type not in {"integer", "number"}:
            return failure(
                "NON_NUMERIC_FIELD",
                f"{arguments.operation} 只能用于数值字段：{arguments.field}",
            )
    if arguments.operation in {"min", "max"} and arguments.field is not None:
        inferred_type = context["fields"][arguments.field]["inferred_type"]
        if inferred_type in {"mixed", "other"}:
            return failure(
                "NON_COMPARABLE_FIELD",
                f"{arguments.operation} 不能用于混合类型字段：{arguments.field}",
            )

    filtered_result = _apply_filters(
        context["rows"], arguments.filters, context["fields"]
    )
    if not filtered_result["ok"]:
        return filtered_result
    matched_rows = filtered_result["data"]
    groups = _group_rows(matched_rows, arguments.group_by)
    results = []
    has_value = False
    for group_key, rows in groups:
        value = _aggregate_value(rows, arguments.operation, arguments.field)
        has_value = has_value or value is not None
        item = {
            group_field: _json_value(group_key[index])
            for index, group_field in enumerate(arguments.group_by)
        }
        item.update(
            {
                "operation": arguments.operation,
                "field": arguments.field,
                "value": value,
            }
        )
        results.append(item)

    if not matched_rows:
        status = "not_found"
    elif arguments.operation != "count" and not has_value:
        status = "no_record"
    else:
        status = "found"
    warnings = []
    if status == "no_record":
        warnings.append("匹配行存在，但聚合字段没有可计算的非空值")
    response = success(
        {
            "status": status,
            "operation": arguments.operation,
            "field": arguments.field,
            "group_by": arguments.group_by,
            "results": results,
            "matched_rows": len(matched_rows),
        },
        "聚合计算完成" if matched_rows else "过滤条件未匹配到记录",
    )
    evidence_fields = list(arguments.group_by)
    if arguments.field:
        evidence_fields.append(arguments.field)
    response.update(
        {
            "evidence": _evidence(
                context,
                fields=evidence_fields,
                filters=arguments.filters,
                matched_rows=len(matched_rows),
            ),
            "warnings": warnings,
            "result_summary": _aggregate_summary(arguments, status, results),
        }
    )
    return response


def load_table_data(file_id: str, sheet_name: str) -> dict[str, Any]:
    """Load rows according to one persisted Schema for trusted Python services."""
    record = database.get_file_record_by_id(file_id)
    if record is None:
        return failure("FILE_NOT_FOUND", "未找到指定 file_id 的文件")
    if record["file_type"] not in {"excel", "csv"}:
        return failure("UNSUPPORTED_FILE_TYPE", "通用表格 Tool 仅支持 Excel 或 CSV")
    if not bool(record["queryable"]):
        return failure("FILE_NOT_QUERYABLE", "文件当前不可查询，请先完成可靠解析")
    schemas = database.get_table_schema_records(file_id, sheet_name)
    if not schemas:
        return failure("SHEET_NOT_FOUND", f"不存在已解析 Sheet：{sheet_name}")
    schema = schemas[0]
    if schema["detection_status"] != "detected":
        return failure("SHEET_NOT_QUERYABLE", f"Sheet 尚未可靠识别：{sheet_name}")
    try:
        path = resolve_by_file_id(file_id)
    except FileLocatorError as exc:
        return failure("FILE_NOT_FOUND", str(exc))

    fields = {field["source_name"]: dict(field) for field in schema["fields"]}
    max_column = max(field["source_index"] for field in schema["fields"])
    workbook = None
    try:
        if record["file_type"] == "csv":
            table = parse_csv_table(path)
            rows = []
            for row_number, values in enumerate(table.rows, start=2):
                row = {
                    field["source_name"]: values[field["source_index"] - 1]
                    if field["source_index"] - 1 < len(values) else None
                    for field in schema["fields"]
                }
                row["__row_number__"] = row_number
                rows.append(row)
        else:
            workbook = load_workbook(
                path,
                read_only=True,
                data_only=True,
                keep_vba=False,
                keep_links=False,
            )
            if sheet_name not in workbook.sheetnames:
                return failure("SHEET_NOT_FOUND", f"Excel 中不存在 Sheet：{sheet_name}")
            worksheet = workbook[sheet_name]
            rows = []
            for row_number, values in enumerate(
                worksheet.iter_rows(
                    min_row=schema["data_start_row"],
                    max_col=max_column,
                    values_only=True,
                ),
                start=schema["data_start_row"],
            ):
                row = {
                    field["source_name"]: values[field["source_index"] - 1]
                    if field["source_index"] - 1 < len(values)
                    else None
                    for field in schema["fields"]
                }
                if any(not _is_null(value) for value in row.values()):
                    row["__row_number__"] = row_number
                    rows.append(row)
    except Exception:
        return failure("TABLE_READ_ERROR", "读取 Excel 表格失败")
    finally:
        if workbook is not None:
            workbook.close()
    return success(
        {
            "record": record,
            "schema": schema,
            "fields": fields,
            "rows": rows,
        },
        "表格读取成功",
    )


def _structured_evidence(
    context: dict[str, Any],
    rows: list[dict[str, Any]],
    selected_fields: list[str],
) -> list[dict[str, Any]]:
    semantic_id = next(
        (
            field["source_name"]
            for field in context["fields"].values()
            if field.get("canonical_name") == "student_id"
        ),
        None,
    )
    evidence = []
    for row in rows:
        row_number = row.get("__row_number__")
        record_key = (
            str(row.get(semantic_id)).strip()
            if semantic_id and row.get(semantic_id) is not None
            else f"row:{row_number}"
        )
        for field in selected_fields:
            value = _json_value(row.get(field))
            source_index = int(context["fields"][field]["source_index"])
            cell = f"{get_column_letter(source_index)}{row_number}"
            table = context["schema"]["sheet_name"]
            evidence.append(
                build_evidence(
                    evidence_id=make_evidence_id(
                        file_id=context["record"]["file_id"],
                        sheet=context["schema"]["sheet_name"],
                        row=row_number,
                        field=field,
                        table=table,
                        cell=cell,
                        value=value,
                    ),
                    source_type="structured",
                    locator_type=(
                        "csv" if context["record"]["file_type"] == "csv" else "excel"
                    ),
                    file_id=context["record"]["file_id"],
                    file_name=context["record"]["file_name"],
                    sheet=context["schema"]["sheet_name"],
                    page_no=None,
                    chunk_id=None,
                    table=table,
                    cell=cell,
                    row_number=(
                        int(row_number)
                        if context["record"]["file_type"] == "csv" else None
                    ),
                    column_name=(
                        field if context["record"]["file_type"] == "csv" else None
                    ),
                    confidence=float(context["schema"]["confidence"]),
                    field=field,
                    record_key=record_key,
                    value_summary=str(value) if value is not None else "null",
                )
            )
    return evidence


def _validate_fields(
    fields: dict[str, dict[str, Any]],
    requested: list[str],
) -> dict[str, Any] | None:
    missing = list(dict.fromkeys(field for field in requested if field not in fields))
    if missing:
        return failure(
            "FIELD_NOT_FOUND",
            "Schema 中不存在字段：" + "、".join(missing),
            {"missing_fields": missing, "available_fields": list(fields)},
        )
    return None


def _apply_filters(
    rows: list[dict[str, Any]],
    filters: list[TableFilter],
    fields: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    prepared = []
    for condition in filters:
        try:
            expected = _coerce_filter_value(
                condition.value,
                fields[condition.field]["inferred_type"],
                condition.op,
            )
        except (TypeError, ValueError):
            return failure(
                "INVALID_FILTER_VALUE",
                f"字段 {condition.field} 的过滤值类型不正确",
            )
        prepared.append((condition, expected))

    matched = []
    for row in rows:
        if all(
            _matches(row.get(condition.field), condition.op, expected)
            for condition, expected in prepared
        ):
            matched.append(row)
    return success(matched, "过滤完成")


def _coerce_filter_value(value: Any, inferred_type: str, op: str) -> Any:
    if op in {"is_null", "not_null", "contains"}:
        return value
    if inferred_type == "integer":
        if isinstance(value, bool):
            raise TypeError
        number = int(value)
        if isinstance(value, float) and not value.is_integer():
            raise ValueError
        return number
    if inferred_type == "number":
        if isinstance(value, bool):
            raise TypeError
        return float(value)
    if inferred_type == "boolean":
        if isinstance(value, bool):
            return value
        if isinstance(value, str) and value.lower() in {"true", "false"}:
            return value.lower() == "true"
        raise ValueError
    if inferred_type == "date":
        return date.fromisoformat(str(value))
    if inferred_type == "datetime":
        return datetime.fromisoformat(str(value))
    if inferred_type == "time":
        return time.fromisoformat(str(value))
    if inferred_type == "string":
        return str(value)
    return value


def _matches(actual: Any, op: str, expected: Any) -> bool:
    is_null = _is_null(actual)
    if op == "is_null":
        return is_null
    if op == "not_null":
        return not is_null
    if is_null:
        return False
    if op == "contains":
        return str(expected) in str(actual)
    if op == "=":
        return actual == expected
    if op == "!=":
        return actual != expected
    try:
        if op == ">":
            return actual > expected
        if op == ">=":
            return actual >= expected
        if op == "<":
            return actual < expected
        if op == "<=":
            return actual <= expected
    except TypeError:
        return False
    return False


def _normalize_orders(
    order_by: TableOrder | list[TableOrder] | None,
) -> list[TableOrder]:
    if order_by is None:
        return []
    return order_by if isinstance(order_by, list) else [order_by]


def _apply_order(
    rows: list[dict[str, Any]],
    orders: list[TableOrder],
) -> list[dict[str, Any]]:
    ordered = list(rows)
    for order in reversed(orders):
        non_null = [row for row in ordered if not _is_null(row.get(order.field))]
        null_rows = [row for row in ordered if _is_null(row.get(order.field))]
        non_null.sort(
            key=lambda row: _sort_key(row.get(order.field)),
            reverse=order.direction == "desc",
        )
        ordered = non_null + null_rows
    return ordered


def _sort_key(value: Any) -> tuple[int, Any]:
    if isinstance(value, bool):
        return 0, int(value)
    if isinstance(value, (int, float)):
        return 1, value
    if isinstance(value, (datetime, date, time)):
        return 2, value.isoformat()
    return 3, str(value)


def _group_rows(
    rows: list[dict[str, Any]],
    group_by: list[str],
) -> list[tuple[tuple[Any, ...], list[dict[str, Any]]]]:
    if not group_by:
        return [((), rows)]
    groups: dict[tuple[Any, ...], list[dict[str, Any]]] = {}
    for row in rows:
        key = tuple(row.get(field) for field in group_by)
        groups.setdefault(key, []).append(row)
    return list(groups.items())


def _aggregate_value(
    rows: list[dict[str, Any]],
    operation: str,
    field: str | None,
) -> int | float | str | None:
    if operation == "count":
        if field is None:
            return len(rows)
        return sum(not _is_null(row.get(field)) for row in rows)
    values = [row.get(field) for row in rows if not _is_null(row.get(field))]
    if not values:
        return None
    if operation == "sum":
        return sum(values)
    if operation == "avg":
        return round(sum(values) / len(values), 6)
    if operation == "min":
        return _json_value(min(values))
    if operation == "max":
        return _json_value(max(values))
    return None


def _evidence(
    context: dict[str, Any],
    *,
    fields: list[str],
    filters: list[TableFilter],
    matched_rows: int,
) -> dict[str, Any]:
    return {
        "file_id": context["record"]["file_id"],
        "file_name": context["record"]["file_name"],
        "sheet": context["schema"]["sheet_name"],
        "fields": fields,
        "filters": [item.model_dump() for item in filters],
        "matched_rows": matched_rows,
    }


def _query_warnings(
    fields: dict[str, dict[str, Any]],
    selected: list[str],
    matched_rows: int,
    returned_rows: int,
) -> list[str]:
    warnings = []
    sensitive = [field for field in selected if bool(fields[field]["sensitive"])]
    if sensitive:
        warnings.append("结果包含可能敏感字段：" + "、".join(sensitive))
    if returned_rows < matched_rows:
        warnings.append("结果已按 limit 截断")
    return warnings


def _aggregate_summary(
    arguments: AggregateTableArguments,
    status: str,
    results: list[dict[str, Any]],
) -> str:
    if status == "not_found":
        return "过滤条件未匹配到记录"
    if status == "no_record":
        return "匹配记录中没有可聚合的非空值"
    return f"完成 {arguments.operation} 计算，返回 {len(results)} 个结果"


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def _is_null(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _validation_message(exc: ValidationError) -> str:
    first = exc.errors()[0]
    location = ".".join(str(item) for item in first.get("loc", []))
    return f"参数 {location or 'root'} 无效：{first.get('msg', '校验失败')}"
