"""Tools for reading and safely appending to existing Word documents."""

import os
import shutil
import tempfile
from pathlib import Path
from typing import Any

from docx import Document

from backend.tools.excel_utils import failure, resolve_data_file, success


def _word_path(file_name: str) -> tuple[Path | None, dict[str, Any] | None]:
    try:
        path = resolve_data_file(file_name, {".docx"})
    except ValueError as exc:
        return None, failure("INVALID_FILE_NAME", str(exc))
    if not path.is_file():
        return None, failure("FILE_NOT_FOUND", f"文件不存在：data/{path.name}")
    return path, None


def read_word(file_name: str) -> dict[str, Any]:
    """Read non-empty paragraphs from an existing Word document."""
    path, error = _word_path(file_name)
    if error is not None:
        return error

    try:
        document = Document(path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    except Exception as exc:
        return failure("WORD_READ_ERROR", f"读取 Word 失败：{exc}")

    return success(
        {"file_name": path.name, "text": "\n".join(paragraphs), "paragraphs": paragraphs},
        "Word 文档读取成功",
    )


def write_word(file_name: str, content: str) -> dict[str, Any]:
    """Safely append content by validating and atomically replacing a temporary copy."""
    path, error = _word_path(file_name)
    if error is not None:
        return error
    if not isinstance(content, str) or not content.strip():
        return failure("INVALID_CONTENT", "追加内容不能为空")

    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=path.parent,
            prefix=f".{path.stem}_",
            suffix=".docx",
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        shutil.copy2(path, temporary_path)
        document = Document(temporary_path)
        document.add_paragraph(content)
        document.save(temporary_path)

        # Reopen the completed copy before replacing the original document.
        Document(temporary_path)
        os.replace(temporary_path, path)
        temporary_path = None
    except Exception as exc:
        return failure("WORD_WRITE_ERROR", f"追加 Word 内容失败，原文件保持不变：{exc}")
    finally:
        if temporary_path is not None:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass

    return success(
        {"file_name": path.name, "path": f"data/{path.name}", "mode": "append"},
        "内容已安全追加到 Word 文档",
    )
