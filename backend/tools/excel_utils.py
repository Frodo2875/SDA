"""Shared helpers for reading project Excel files safely."""

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = PROJECT_ROOT / "data"


def success(data: Any, message: str) -> dict[str, Any]:
    """Build a successful tool response."""
    return {"ok": True, "data": data, "error_code": None, "message": message}


def failure(error_code: str, message: str, data: Any = None) -> dict[str, Any]:
    """Build a failed tool response."""
    return {
        "ok": False,
        "data": data,
        "error_code": error_code,
        "message": message,
    }


def resolve_data_file(file_name: str, suffixes: set[str]) -> Path:
    """Resolve a plain file name inside data/ and reject path traversal."""
    if not isinstance(file_name, str) or not file_name.strip():
        raise ValueError("文件名不能为空")

    clean_name = file_name.strip()
    if Path(clean_name).name != clean_name:
        raise ValueError("文件名不能包含目录路径")

    path = DATA_DIR / clean_name
    if path.suffix.lower() not in suffixes:
        raise ValueError(f"不支持的文件类型：{path.suffix or '无扩展名'}")
    return path


def normalize_student_id(value: Any) -> str:
    """Convert a worksheet student ID to a normalized string."""
    if value is None:
        return ""
    return str(value).strip()


def read_excel_rows(file_name: str, required_fields: list[str]) -> dict[str, Any]:
    """Read the active worksheet as dictionaries after validating its fields."""
    try:
        path = resolve_data_file(file_name, {".xlsx"})
    except ValueError as exc:
        return failure("INVALID_FILE_NAME", str(exc))

    if not path.is_file():
        return failure("FILE_NOT_FOUND", f"文件不存在：data/{path.name}")

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            values = worksheet.iter_rows(values_only=True)
            header_row = next(values, None)
            if header_row is None:
                return failure("EMPTY_WORKBOOK", f"Excel 文件为空：data/{path.name}")

            headers = [str(value).strip() if value is not None else "" for value in header_row]
            missing_fields = [field for field in required_fields if field not in headers]
            if missing_fields:
                return failure(
                    "MISSING_FIELDS",
                    f"Excel 缺少必要字段：{', '.join(missing_fields)}",
                    {"missing_fields": missing_fields},
                )

            rows = []
            for values_row in values:
                padded_row = tuple(values_row) + (None,) * (len(headers) - len(values_row))
                row = {header: padded_row[index] for index, header in enumerate(headers) if header}
                if not any(value is not None for value in row.values()):
                    continue
                if "学号" in row:
                    row["学号"] = normalize_student_id(row["学号"])
                rows.append(row)
        finally:
            workbook.close()
    except Exception as exc:  # openpyxl exposes several format/IO exception types
        return failure("EXCEL_READ_ERROR", f"读取 Excel 失败：{exc}")

    return success(rows, f"成功读取 data/{path.name}")
